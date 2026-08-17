#!/usr/bin/env python3
"""CONTROL Steps 6–7: internal MD summaries + client Word from feature indexes."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x5A, 0x88)
HEADER_BG = "1B3A5F"
ALT_ROW = "F3F6FA"
MUTED = RGBColor(0x55, 0x55, 0x55)

PROJECTS = {
    "CLAIM": {
        "root": Path(__file__).resolve().parent / "ClaimManagement",
        "code": "CLAIM",
        "name": "Claim Management System (KPID / SPLOCA)",
        "short": "Claim Management System",
        "assumptions": [
            "Feature IDs match the Claim Management Statement of Work.",
            "Access is Person mapped to Post, with Roles on the Post. This identity is also used by Legal Management and GIS field valuation.",
            "FE dossiers remain in the existing FE application; this estimate covers import of FE details, Directors and Promoters, and Investment Schemes only.",
            "In-Karnataka intake is staff-mediated from originals; out-of-Karnataka uses a constrained path. Submit requires the full mandatory document pack.",
            "GIS field valuation and Legal Management are separate estimates and are not included in these hours.",
            "Hours are indicative person-hours for planning and commercial discussion. Hosting, AMC, scanning bureau and classroom training are not included unless separately ordered.",
        ],
        "exclusions": [
            "FE Management dossier rebuild",
            "Legal Management Application",
            "GIS field valuation (eValueI / eValueM)",
            "Generative AI",
            "Hardware, hosting AMC, scanning bureau",
            "Market intelligence and police investigation process",
        ],
        "phasing": [
            ("P1 — Identity and FE import", "01, 02"),
            ("P2 — Recovery feed", "03"),
            ("P3 — Depositor, investment, intake", "04, 05, 06"),
            ("P4 — Verification and settlement", "07, 08"),
            ("P5 — Reports and Special Court outputs", "09"),
        ],
    },
    "EVAL": {
        "root": Path(__file__).resolve().parent / "Valuation",
        "code": "EVAL",
        "name": "GIS Field Valuation — eValueI / eValueM (KPID / SPLOCA)",
        "short": "GIS Field Valuation (eValueI / eValueM)",
        "assumptions": [
            "Feature IDs match the GIS Field Valuation Statement of Work.",
            "User management is consumed from Claim Management Module 01; these hours cover valuation wiring, Posts, assignment rights and FAQ — not a second identity product.",
            "Auction sale and recovered-fund deposit remain in Claim Management; this estimate includes release of the accepted valuation pack and auction-display photographs.",
            "Form V-5 and office IT inventory are out of scope.",
            "Hours are indicative person-hours for planning and commercial discussion. Hosting, AMC and Play Store fees (if any) are not included unless separately ordered.",
        ],
        "exclusions": [
            "Rebuild of user management (Claim Management Module 01)",
            "Claim settlement, auction sale and recovered-fund deposit",
            "Legal Management Application",
            "Form V-5",
            "Office IT inventory and maintenance tickets",
            "Generative AI",
            "Hardware and hosting AMC",
        ],
        "phasing": [
            ("P1 — Identity and masters", "01, 02"),
            ("P2 — Office asset register", "03"),
            ("P3 — Field capture and methods", "04, 05"),
            ("P4 — Movables and cash", "06"),
            ("P5 — CA, billing, Android", "07, 08, 09"),
        ],
    },
    "LEGAL": {
        "root": Path(__file__).resolve().parent / "LegalManagement",
        "code": "LEGAL",
        "name": "Legal Management Application (KPID / SPLOCA)",
        "short": "Legal Management Application",
        "assumptions": [
            "Feature IDs match the Legal Management Statement of Work.",
            "User management is consumed from Claim Management Module 01; these hours cover Legal wiring, Limited/Community spaces and document rights — not a second identity product.",
            "Generative AI is not used for legal drafting. Templates are stored and human-verified.",
            "Hours are indicative person-hours for planning and commercial discussion. Hosting, AMC, classroom training and Play Store fees (if any) are not included unless separately ordered.",
        ],
        "exclusions": [
            "Rebuild of user management (Claim Management Module 01)",
            "Claim settlement",
            "GIS field valuation (eValueI / eValueM)",
            "FE Management dossier",
            "Market intelligence and investigation process",
            "Generative AI",
            "Hardware and hosting AMC",
        ],
        "phasing": [
            ("P1 — Identity and masters", "01, 02"),
            ("P2 — Library and search", "03"),
            ("P3 — Matters and work", "04, 05, 06, 07"),
            ("P4 — Knowledge", "08"),
            ("P5 — Android", "09"),
        ],
    },
}


def parse_index(path: Path) -> tuple[str, list[dict], int]:
    text = path.read_text(encoding="utf-8")
    heading = re.search(r"^##\s+(.+)$", text, re.M)
    title = heading.group(1).strip() if heading else path.parent.parent.name
    rows = []
    for line in text.splitlines():
        if not re.match(r"^\|\s*F\d{2}-\d{3}\s*\|", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 3:
            continue
        hours = int(re.sub(r"[^\d]", "", cells[2]) or "0")
        rows.append(
            {
                "id": cells[0],
                "name": cells[1],
                "hours": hours,
                "notes": cells[3] if len(cells) > 3 else "",
            }
        )
    total_m = re.search(r"\*\*(\d[\d,]*)\s*hours?\*\*", text, re.I)
    total = int(total_m.group(1).replace(",", "")) if total_m else sum(r["hours"] for r in rows)
    return title, rows, total


def load_project(cfg: dict) -> dict:
    root: Path = cfg["root"]
    registry = json.loads((root / f"{cfg['code']}_Registry.json").read_text(encoding="utf-8"))
    modules = []
    for m in registry["modules"]:
        idx = root / m["folder"] / "features" / "_index.md"
        title, rows, total = parse_index(idx)
        if total != m["total_hours"] or len(rows) != m["feature_count"]:
            raise SystemExit(
                f"{cfg['code']} {m['nn']}: index {len(rows)}/{total} vs registry "
                f"{m['feature_count']}/{m['total_hours']}"
            )
        modules.append(
            {
                "nn": m["nn"],
                "name": m["name"],
                "type": m["type"],
                "features": rows,
                "hours": total,
                "count": len(rows),
            }
        )
    grand = sum(x["hours"] for x in modules)
    fcount = sum(x["count"] for x in modules)
    return {"cfg": cfg, "modules": modules, "grand": grand, "fcount": fcount}


def write_internal_md(data: dict) -> None:
    cfg, modules, grand, fcount = data["cfg"], data["modules"], data["grand"], data["fcount"]
    root: Path = cfg["root"]
    code = cfg["code"]
    go = grand
    opt = 0

    lines = [
        f"# {code} — Estimation Summary",
        "",
        f"> **Detailed module & feature-wise hours:** [`{code}_Estimation_Summary_FeatureWise.md`]({code}_Estimation_Summary_FeatureWise.md)",
        "",
        f"Source: `SoW_{code}_Detailed.md`",
        f"Modules file: `{code}_Modules.md`",
        "Internal roll-up from module feature indexes. Hours from Estimation Master.",
        "",
        "---",
        "",
        "## Totals",
        "",
        f"- **Modules:** {len(modules)}",
        f"- **Features:** {fcount}",
        f"- **Grand total:** **{grand:,} hours**",
        f"- **Recommended go-live:** **{go:,} hours**",
        f"- **Optional / conditional:** **{opt} hours**",
        "",
        "## Per-module totals",
        "",
        "| # | Module | Type | Features | Hours |",
        "|---|--------|------|---------:|------:|",
    ]
    for m in modules:
        lines.append(f"| {m['nn']} | {m['name']} | {m['type']} | {m['count']} | {m['hours']:,} |")
    lines.append(f"| | **Total** | | **{fcount}** | **{grand:,}** |")
    lines += [
        "",
        "## Go-live vs optional",
        "",
        f"- Go-live (all modules 01–09): **{go:,} hours**",
        f"- Optional in this baseline: **{opt} hours**",
        "",
        "## Assumptions / risks carried forward",
        "",
    ]
    for i, a in enumerate(cfg["assumptions"], 1):
        lines.append(f"{i}. {a}")
    lines += [
        "",
        "## Suggested next actions",
        "",
        "1. Review hours with delivery and commercial leads.",
        "2. Confirm shared identity (Claim Module 01) availability for Legal and Valuation.",
        "3. Share the client Word estimate for RFP / commercial discussion.",
        "",
        "---",
        "",
        "*End of internal roll-up*",
        "",
    ]
    (root / f"{code}_Estimation_Summary.md").write_text("\n".join(lines), encoding="utf-8")

    fw = [
        f"# {code} — Estimation Summary (Feature-Wise)",
        "",
        "Internal detail. Hours from Estimation Master via feature files.",
        "",
        f"**Grand total:** {grand:,} hours · **Go-live:** {go:,} · **Optional:** {opt}",
        "",
        "## Module-wise summary",
        "",
        "| # | Module | Type | Features | Hours | % of go-live |",
        "|---|--------|------|---------:|------:|-------------:|",
    ]
    for m in modules:
        pct = (m["hours"] / go * 100) if go else 0
        fw.append(
            f"| {m['nn']} | {m['name']} | {m['type']} | {m['count']} | {m['hours']:,} | {pct:.1f}% |"
        )
    fw.append(f"| | **Total** | | **{fcount}** | **{grand:,}** | **100%** |")
    for m in modules:
        fw += [
            "",
            f"## {m['nn']} — {m['name']} ({m['hours']:,} hours)",
            "",
            "| Feature ID | Feature Name | Hours | Notes |",
            "|---|---|---:|---|",
        ]
        for r in m["features"]:
            notes = r["notes"].replace("|", "/")
            fw.append(f"| {r['id']} | {r['name']} | {r['hours']} | {notes} |")
    fw += [
        "",
        "## Indicative phasing",
        "",
        "| Phase | Modules |",
        "|-------|---------|",
    ]
    for phase, mods in cfg["phasing"]:
        fw.append(f"| {phase} | {mods} |")
    fw += [
        "",
        "## Exclusions (not in feature hours)",
        "",
    ]
    for e in cfg["exclusions"]:
        fw.append(f"- {e}")
    fw.append("")
    (root / f"{code}_Estimation_Summary_FeatureWise.md").write_text("\n".join(fw), encoding="utf-8")


def set_run_font(run, *, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(p, *, space_before=0, space_after=6, line_spacing=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CBD5E1")
        borders.append(el)
    tblPr.append(borders)


def add_table(doc: Document, rows: list[list[str]]):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = True
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            style_paragraph(p, space_before=2, space_after=2, line_spacing=1.05)
            val = row[c_idx] if c_idx < len(row) else ""
            run = p.add_run(val)
            set_run_font(
                run,
                size=9 if r_idx == 0 else 10,
                bold=r_idx == 0,
                color=RGBColor(0xFF, 0xFF, 0xFF) if r_idx == 0 else None,
            )
            if r_idx == 0:
                set_cell_shading(cell, HEADER_BG)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, ALT_ROW)
    doc.add_paragraph()


def add_page_numbers(doc: Document, footer_label: str):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"SPLOCA | {footer_label}  |  Page ")
    set_run_font(run, size=8, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld1)
    run2._r.append(instr)
    run2._r.append(fld2)
    set_run_font(run2, size=8, color=MUTED)


def write_client_docx(data: dict) -> None:
    cfg, modules, grand, fcount = data["cfg"], data["modules"], data["grand"], data["fcount"]
    root: Path = cfg["root"]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SPECIAL OFFICER & COMPETENT AUTHORITY (SPLOCA)")
    set_run_font(run, size=14, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Government of Karnataka")
    set_run_font(run, size=12, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=18, space_after=6)
    run = p.add_run("Module & Feature-wise Effort Estimate")
    set_run_font(run, size=20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(cfg["short"])
    set_run_font(run, size=16, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Client Copy · Effort in person-hours")
    set_run_font(run, size=12, color=MUTED)
    doc.add_page_break()

    h = doc.add_heading("1. Purpose", level=1)
    style_paragraph(h, space_before=0, space_after=8)
    p = doc.add_paragraph()
    p.add_run(
        "This estimate sets out indicative person-hours to design, build, test and document "
        f"{cfg['short']} as specified in the Statement of Work. Hours cover application software "
        "for the modules listed below. They do not cover hosting infrastructure, annual maintenance, "
        "scanning bureau operations, or classroom training unless those items are ordered separately."
    )

    h = doc.add_heading("2. Summary totals", level=1)
    add_table(
        doc,
        [
            ["Item", "Person-hours"],
            ["Recommended go-live", f"{grand:,}"],
            ["Optional / conditional", "0"],
            ["All-in (this baseline)", f"{grand:,}"],
            ["Features", str(fcount)],
            ["Modules", str(len(modules))],
        ],
    )

    h = doc.add_heading("3. Module-wise summary", level=1)
    rows = [["Module", "Type", "Features", "Hours", "% of go-live"]]
    for m in modules:
        pct = f"{m['hours'] / grand * 100:.1f}%" if grand else "0%"
        rows.append([f"{m['nn']}. {m['name']}", m["type"], str(m["count"]), f"{m['hours']:,}", pct])
    rows.append(["Total", "", str(fcount), f"{grand:,}", "100%"])
    add_table(doc, rows)

    h = doc.add_heading("4. Feature-wise estimate", level=1)
    for m in modules:
        doc.add_heading(f"{m['nn']}. {m['name']} — {m['hours']:,} hours", level=2)
        frows = [["Feature ID", "Feature Name", "Hours", "Notes"]]
        for r in m["features"]:
            frows.append([r["id"], r["name"], str(r["hours"]), r["notes"]])
        add_table(doc, frows)

    type_hours = {"Shared": 0, "Business": 0, "Technical": 0}
    for m in modules:
        type_hours[m["type"]] = type_hours.get(m["type"], 0) + m["hours"]
    all_feats = [(r["hours"], r["id"], r["name"], m["nn"]) for m in modules for r in m["features"]]
    all_feats.sort(reverse=True)

    h = doc.add_heading("5. Effort profile", level=1)
    add_table(
        doc,
        [["Type", "Hours"]]
        + [[k, f"{v:,}"] for k, v in type_hours.items() if v]
        + [["Total", f"{grand:,}"]],
    )
    p = doc.add_paragraph()
    run = p.add_run("Largest features (indicative)")
    set_run_font(run, bold=True, color=NAVY)
    top = [["Feature ID", "Feature Name", "Hours"]]
    for hrs, fid, name, _nn in all_feats[:8]:
        top.append([fid, name, str(hrs)])
    add_table(doc, top)

    h = doc.add_heading("6. Indicative delivery phasing", level=1)
    add_table(doc, [["Phase", "Primary modules"]] + [list(x) for x in cfg["phasing"]])
    p = doc.add_paragraph()
    p.add_run(
        "Phasing is indicative for planning. Exact go-live grouping is a commercial and Authority decision."
    )

    h = doc.add_heading("7. Items not included in feature hours", level=1)
    for e in cfg["exclusions"]:
        doc.add_paragraph(e, style="List Bullet")

    h = doc.add_heading("8. Planning assumptions", level=1)
    for i, a in enumerate(cfg["assumptions"], 1):
        doc.add_paragraph(f"{i}. {a}")

    h = doc.add_heading("9. Disclaimer", level=1)
    p = doc.add_paragraph()
    p.add_run(
        "These figures are indicative person-hours for planning and commercial discussion. "
        "They assume a production-quality build with validation, error handling, logging, security "
        "and unit testing. They are not a fixed-price quotation until commercials, staffing mix "
        "and delivery window are agreed."
    )

    add_page_numbers(doc, f"{cfg['short']} estimate")
    out = root / f"{cfg['code']}_Estimation_Summary_FeatureWise_Client.docx"
    doc.save(out)
    print(f"Wrote {out}")


def main(codes: list[str]) -> None:
    for code in codes:
        cfg = PROJECTS[code]
        data = load_project(cfg)
        write_internal_md(data)
        write_client_docx(data)
        print(f"{code}: {data['fcount']} features, {data['grand']} hours")


if __name__ == "__main__":
    import sys

    codes = sys.argv[1:] or ["CLAIM", "EVAL"]
    main(codes)
