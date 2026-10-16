"""Build BPA estimation artefacts (Steps 2-3, 5, 6, 7) from feature data.

Run from repo root:
    .venv/bin/python Projects/BDA_Afterauction/_pipeline/build_bpa.py
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

from bpa_data_a import FEATURES as FA  # noqa: E402
from bpa_data_b import FEATURES as FB  # noqa: E402
from bpa_data_c import FEATURES as FC  # noqa: E402
from bpa_data_d import FEATURES as FD  # noqa: E402
from bpa_data_e import FEATURES as FE  # noqa: E402

FEATURES = FA + FB + FC + FD + FE
CODE = "BPA"
PROJECT = "BDA Post-Auction Award, Payment and Completion Management Module"
CLIENT = "Bangalore Development Authority (BDA)"
TODAY = dt.date.today().isoformat()

CX = {"S": "Simple", "E": "Easy", "M": "Medium", "C": "Complex", "V": "Very High"}
HOURS = {
    "UI": {"S": 2, "E": 4, "M": 6, "C": 12, "V": 16},
    "API": {"S": 2, "E": 4, "M": 6, "C": 8, "V": 16},
    "Database": {"S": 1, "E": 2, "M": 4, "C": 6, "V": 10},
    "Integration": {"S": 4, "E": 8, "M": 16, "C": 24, "V": 40},
    "Testing": {"S": 2, "E": 4, "M": 6, "C": 8, "V": 12},
    "Documentation": {"S": 1, "E": 2, "M": 3, "C": 4, "V": 6},
}
AREA_KEYS = [("UI", "screens"), ("API", "apis"), ("Database", "db"), ("Integration", "integ"), ("Testing", "tests"),
             ("Documentation", "docs")]
GO_LIVE = "Go-live"

MODULES = [
    dict(nn="01", name="User, Role & Access Management", type="Shared",
         summary="Staff login and sessions, user and role management, configurable permissions with DS-wise isolation, officer DSC registration",
         source="Explicit (SOW §4.10, §4.12)",
         in_scope="Staff authentication with OTP and lockout; user register; role assignment with DS office; permission matrix; DS-wise record isolation; DSC certificate registration and signing authority",
         out_scope="Bidder login (Module 05); single sign-on with other BDA systems",
         actors="Super Admin; Auction Admin; all BDA roles",
         deps="F12-003 (OTP SMS / e-mail); F13-003 (DSC verification); F14-001 (audit)",
         notes="Access model User → Role → Permissions with DS-wise data isolation; bidders use a separate OTP login"),
    dict(nn="02", name="Auction Rules & Configuration", type="Business",
         summary="Instalment and payment rules, DS area allocation, approval hierarchy and thresholds, stage timelines, document templates, reference masters",
         source="Explicit (SOW §4.2, §4.4A, §4.5, §4.6, §4.11)",
         in_scope="Versioned, effective-dated rules and bilingual templates maintained by the Auction Admin",
         out_scope="Notification templates (Module 12)",
         actors="Auction Admin; Super Admin; Approving Authority (template approval)",
         deps="Consumed by Modules 03–09",
         notes="Delivers the 'no code change for rule revisions' requirement"),
    dict(nn="03", name="KPPP Result Intake & Site Package", type="Business",
         summary="KPPP API receipt, controlled report upload, matching and exceptions, CDRMS site package, locking and versioned re-import",
         source="Explicit (SOW §4.1)",
         in_scope="Result intake by API or upload; site matching; exception queue; read-only site package; integrity lock; re-import with approval",
         out_scope="Any live bidding or pre-auction activity on KPPP",
         actors="System; Auction Admin; Result Verifier",
         deps="F13-001; F02-006; KPPP and CDRMS access",
         notes="Upload route ensures scope is not blocked if the KPPP API is unavailable"),
    dict(nn="04", name="Result Verification & Approval", type="Business",
         summary="Result review workbench, single-bidder and abnormal bid flags, multi-level DSC approval, draft proceedings, rejected / unsold intimation",
         source="Explicit (SOW §4.2)",
         in_scope="Verification, flags, configurable approval up to Government, site-wise decisions, proceedings, allotment section intimation",
         out_scope="Re-auction execution on KPPP",
         actors="Result Verifier; Approving Authority; Auction Admin",
         deps="Module 03; F02-003; F02-005; F13-003",
         notes="Workflow embedded; no separate workflow module"),
    dict(nn="05", name="Successful Bidder Notification & Account", type="Business",
         summary="Bidder account creation, OTP login and profile, bilingual DSC-signed intimation letter, dispatch with delivery status",
         source="Explicit (SOW §4.3)",
         in_scope="Account on approval; mobile / e-mail OTP; intimation letter; e-mail and SMS dispatch",
         out_scope="Bidder registration on KPPP",
         actors="System; Successful Bidder; Auction Admin",
         deps="Module 04; F02-005; Module 12; F13-003",
         notes="Bidder login is separate from staff login"),
    dict(nn="06", name="Payment & Instalment Management", type="Business",
         summary="Instalment demands, online and offline payments, daily reconciliation, ledger, reminders, extensions, defaults and disputes",
         source="Explicit (SOW §4.4)",
         in_scope="25% / 50% / 25% instalments with EMD adjustment; gateway payments; offline verification; bank reconciliation; extension, default and dispute handling",
         out_scope="EMD collection and refund (KPPP); accounting system postings",
         actors="Successful Bidder; Payment Reconciler; Finance Officer; Approving Authority; Auction Admin",
         deps="F02-001; Module 12; e-payment gateway and bank settlement files",
         notes="Highest-risk financial module; gateway connector is built here"),
    dict(nn="07", name="Case Forwarding & Deputy Secretary Processing", type="Business",
         summary="Automatic forwarding by area on 75% payment, re-assignment, case file workspace, processing and document requests, timelines and alerts",
         source="Explicit (SOW §4.5)",
         in_scope="DS1 / DS2 / DS3 queues; case file; notes; bidder document requests and uploads; stage timers",
         out_scope="Registration of the deed",
         actors="Deputy Secretary; Case Worker; Auction Admin; Successful Bidder",
         deps="F02-002; F02-004; F06-005; Module 03 (site package)",
         notes="DS-wise isolation from F01-003 applies"),
    dict(nn="08", name="Draft Sale Deed", type="Business",
         summary="Template-based draft sale deed, review and approval with versions and DSC, bidder access, final instalment trigger",
         source="Explicit (SOW §4.6)",
         in_scope="Generation, editing, approval, signing, bidder download, final 25% demand",
         out_scope="Stamp duty computation and registration",
         actors="Case Worker; Deputy Secretary; Successful Bidder",
         deps="F02-005; F13-003; F06-001",
         notes=""),
    dict(nn="09", name="Registration Recording & Completion", type="Business",
         summary="Registration reference and deed copy, Kaveri exchange where permitted, completion record, status update to allotment section and CDRMS",
         source="Explicit (SOW §4.7)",
         in_scope="Registration capture; conditional Kaveri link; completion; CDRMS write-back",
         out_scope="Registration itself (Department of Stamps and Registration)",
         actors="Case Worker; Deputy Secretary; Successful Bidder; System",
         deps="Module 06 (100% payment); F13-001; CDRMS update access",
         notes="Kaveri exchange is conditional on departmental permission"),
    dict(nn="10", name="Bidder Tracking Portal", type="Business",
         summary="My awarded sites, site details and eleven-stage tracker, payments and receipts, documents, notifications and profile, bilingual GIGW portal",
         source="Explicit (SOW §4.8)",
         in_scope="Self-service tracking for successful bidders in Kannada and English",
         out_scope="Public auction information portal",
         actors="Successful Bidder",
         deps="Modules 05–09; F13-004",
         notes="Payment initiation and uploads reuse Module 06 and 07 operations"),
    dict(nn="11", name="Dashboards, MIS & Reports", type="Business",
         summary="Commissioner, DS and Finance dashboards, standard MIS reports, configurable report builder",
         source="Explicit (SOW §4.9)",
         in_scope="Role dashboards; auction-, layout-, DS- and period-wise reports in Excel / PDF; report builder",
         out_scope="External BI tools",
         actors="Commissioner / Read-only Viewer; Deputy Secretary; Finance Officer; Auction Admin",
         deps="All business modules; F01-003",
         notes="DS-wise isolation applies to all outputs"),
    dict(nn="12", name="Notification Service", type="Shared",
         summary="Bilingual DLT templates, stage event engine, e-mail / SMS / WhatsApp connectors, delivery log",
         source="Explicit (SOW §4.11)",
         in_scope="Event and schedule-based notifications for every stage; delivery tracking",
         out_scope="Gateway subscription charges",
         actors="Auction Admin; System; recipients",
         deps="BDA e-mail, SMS (DLT) and WhatsApp accounts; F13-001",
         notes="Used by Modules 01 and 03–10"),
    dict(nn="13", name="Integration Hub", type="Technical",
         summary="Connector framework and credentials, integration monitoring and error queue, DSC / e-Sign service, Web GIS map",
         source="Derived (SOW §4.13)",
         in_scope="Shared integration services; signing; maps",
         out_scope="KPPP, CDRMS, payment and Kaveri business mapping (built in owning features)",
         actors="Super Admin; Auction Admin; System",
         deps="BDA e-Sign / DSC arrangement; BDA Web GIS service",
         notes="Business connectors live in Modules 03, 06 and 09 and reuse this framework"),
    dict(nn="14", name="Audit, Security & Compliance", type="Technical",
         summary="Hash-chained audit trail, audit viewer and Section 63 support, security controls, DPDP compliance",
         source="Explicit (SOW §4.12)",
         in_scope="Tamper-evident audit; integrity verification; encryption; OWASP controls; privacy and data principal requests",
         out_scope="VAPT execution by the CERT-In empanelled auditor",
         actors="Super Admin; auditors; Successful Bidder (privacy requests)",
         deps="NTP source; hosting key management",
         notes="Audit is read-only for every role including Super Admin"),
    dict(nn="15", name="Platform Operations & Support", type="Technical",
         summary="Environments and releases, backup and DR, monitoring and SLA reporting, support helpdesk",
         source="Derived (SOW §4.14, §4.17, §8)",
         in_scope="Dev / UAT / production; pipelines; backups within RPO / RTO; availability monitoring; tickets and monthly report",
         out_scope="Hosting infrastructure and licences",
         actors="Super Admin; support team; BDA IT",
         deps="State Data Centre or BDA infrastructure",
         notes="Supports SLA measurement and monthly reporting"),
    dict(nn="16", name="Legacy Case Migration", type="Technical",
         summary="Migration templates and staging, validation and reconciliation, activation of migrated cases at the correct stage",
         source="Explicit (SOW §4.15)",
         in_scope="Pending cases from earlier auctions with payments and documents",
         out_scope="Data entry of paper records (BDA provides data in templates)",
         actors="Auction Admin; Finance Officer; Deputy Secretary; migration team",
         deps="Modules 03–09; legacy data from BDA",
         notes="BDA sign-off required before go-live"),
]


def slugify(text: str) -> str:
    text = text.lower().replace("&", "and")
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


for m in MODULES:
    m["slug"] = slugify(m["name"])
    m["folder"] = f"modules/{m['nn']}-{m['slug']}"


def rows(f):
    out = []
    for area, key in AREA_KEYS:
        for name, cx in f.get(key, []):
            out.append((area, name, CX[cx], HOURS[area][cx]))
    return out


def area_totals(f):
    t = {a: 0 for a, _ in AREA_KEYS}
    for area, _, _, h in rows(f):
        t[area] += h
    return t


for f in FEATURES:
    f.setdefault("notes", GO_LIVE)
    f["client"] = "Recommended for go-live" if f["notes"] == GO_LIVE else "Conditional — subject to departmental permission"
    f["slug"] = slugify(f["name"])
    f["nn"] = f["id"][1:3]
    f["seq"] = f["id"][4:]
    f["hours"] = sum(r[3] for r in rows(f))
    f["areas"] = area_totals(f)


def module_of(f):
    return next(m for m in MODULES if m["nn"] == f["nn"])


def feats(m):
    return [f for f in FEATURES if f["nn"] == m["nn"]]


def bullets(items):
    return "\n".join(f"- {i}" for i in items) if items else "- None"


# ---------------------------------------------------------------- validation
REQUIRED = ("desc", "obj", "actors", "journey", "caps", "screens", "apis", "db", "tests", "docs", "inputs", "outputs",
            "rules", "valid", "calc", "flow", "search", "documents", "reports", "notif", "sec", "ints", "deps", "assume",
            "risks", "nfr", "est")


def validate():
    errors = []
    sow = (ROOT / f"SoW_{CODE}_Detailed.md").read_text()
    sow_feats = dict(re.findall(r"^\| (F\d{2}-\d{3}) \| ([^|]+?) \|", sow, re.M))
    ids = [f["id"] for f in FEATURES]
    if len(ids) != len(set(ids)):
        errors.append("duplicate feature IDs")
    for f in FEATURES:
        if f["id"] not in sow_feats:
            errors.append(f"{f['id']} missing in SOW")
        elif sow_feats[f["id"]].strip() != f["name"]:
            errors.append(f"{f['id']} name mismatch: SOW '{sow_feats[f['id']].strip()}' vs '{f['name']}'")
        if len(f.get("journey", [])) != 8:
            errors.append(f"{f['id']} journey has {len(f.get('journey', []))} points")
        for k in REQUIRED:
            if k not in f:
                errors.append(f"{f['id']} missing {k}")
        for area, key in AREA_KEYS:
            for name, cx in f.get(key, []):
                if cx not in CX:
                    errors.append(f"{f['id']} bad complexity {cx} on {name}")
    for fid in sow_feats:
        if fid not in ids:
            errors.append(f"{fid} in SOW but no feature data")
    for m in MODULES:
        if f"### {m['nn']}. {m['name']}" not in sow:
            errors.append(f"module heading missing in SOW: {m['nn']}. {m['name']}")
    if errors:
        raise SystemExit("Validation failed:\n" + "\n".join(errors))


# ---------------------------------------------------------------- Step 5
def render_feature(f):
    m = module_of(f)
    lines = [
        f"# {f['id']} — {f['name']}", "",
        "## Feature Specification", "",
        "| Field | Value |", "|-------|-------|",
        f"| Feature ID | {f['id']} |",
        f"| Feature Name | {f['name']} |",
        f"| Module | {m['nn']} — {m['name']} |",
        f"| Module Type | {m['type']} |",
        f"| Release | {f['notes']} |", "",
        "### Description", f["desc"], "",
        "### Business Objective", f["obj"], "",
        "### Actors", bullets(f["actors"]), "",
        "### User Journey",
    ]
    labels = ["Who creates", "How they create", "How they edit", "How they send for approval",
              "How the approver sees pending items", "How the approver decides",
              "How the applicant / maker sees history", "How the approver / auditor sees history"]
    for i, (lab, txt) in enumerate(zip(labels, f["journey"]), 1):
        lines.append(f"{i}. **{lab}:** {txt}")
    lines += ["", "### Functional Capabilities", bullets(f["caps"]), ""]
    lines += ["### UI Screens", bullets([s for s, _ in f["screens"]]) if f["screens"] else "- None (no user interface in this feature)", ""]
    lines += ["### Logical API Operations", bullets([a for a, _ in f["apis"]]) if f["apis"] else "- None", ""]
    for title, key in [("Inputs", "inputs"), ("Outputs", "outputs"), ("Business Rules", "rules"),
                       ("Validations", "valid"), ("Calculations", "calc")]:
        lines += [f"### {title}", bullets(f[key]), ""]
    lines += ["### Workflow", f["flow"], ""]
    for title, key in [("Search Requirements", "search"), ("Documents", "documents"), ("Reports", "reports"),
                       ("Notifications", "notif"), ("Security", "sec"), ("Integrations", "ints"),
                       ("Dependencies", "deps"), ("Assumptions", "assume"), ("Risks", "risks"),
                       ("Non Functional Requirements", "nfr")]:
        lines += [f"### {title}", bullets(f[key]), ""]
    lines += ["## Estimation Breakdown", "", "| Area | Artefact | Complexity | Hours |", "|------|----------|------------|------:|"]
    for area, name, cx, h in rows(f):
        lines.append(f"| {area} | {name} | {cx} | {h} |")
    lines += ["", f"**Feature Total: {f['hours']} hours**", ""]
    a = f["areas"]
    lines += ["| UI | API | Database | Integration | Testing | Documentation | Total |",
              "|---:|---:|---:|---:|---:|---:|---:|",
              f"| {a['UI']} | {a['API']} | {a['Database']} | {a['Integration']} | {a['Testing']} | {a['Documentation']} | **{f['hours']}** |",
              "", "### Estimation Assumptions",
              "- Every UI screen and logical API operation listed above is estimated as its own row using the Estimation Master.",
              "- Enterprise-grade implementation: validation, error handling, logging, audit hooks, security checks and unit tests included in each row."]
    lines += [f"- {e}" for e in f["est"]]
    if not f.get("integ"):
        lines.append("- No external integration channel in this feature; Integration row omitted.")
    return "\n".join(lines) + "\n"


def write_features():
    for m in MODULES:
        fdir = ROOT / m["folder"] / "features"
        fdir.mkdir(parents=True, exist_ok=True)
        for old in fdir.glob("F*.md"):
            old.unlink()
        ms = feats(m)
        for f in ms:
            f["file"] = f"{m['folder']}/features/{f['id']}-{f['slug']}.md"
            (ROOT / f["file"]).write_text(render_feature(f))
        m["hours"] = sum(f["hours"] for f in ms)
        idx = ["# Feature Index", "", f"## {m['nn']} — {m['name']}", "",
               "| Feature ID | Feature Name | Total Hours | Notes |", "|------------|--------------|------------:|-------|"]
        idx += [f"| {f['id']} | {f['name']} | {f['hours']} | {f['notes']} |" for f in ms]
        idx += ["", f"**Module Total:** **{m['hours']} hours**", ""]
        (fdir / "_index.md").write_text("\n".join(idx))
        readme = [f"# {m['nn']}. {m['name']}", "", f"Type: {m['type']}", f"Summary: {m['summary']}.", "",
                  f"- **Source:** {m['source']}", f"- **In Scope:** {m['in_scope']}", f"- **Out of Scope:** {m['out_scope']}",
                  f"- **Key Actors:** {m['actors']}", f"- **Dependencies:** {m['deps']}",
                  f"- **Notes:** {m['notes'] or '—'}", "",
                  f"Features: {len(ms)} · **{m['hours']} hours** (see `features/_index.md`)", ""]
        (ROOT / m["folder"] / "README.md").write_text("\n".join(readme))


# ---------------------------------------------------------------- Steps 2-3
def write_modules_md():
    out = [f"# {PROJECT} — Module List", "", f"Source: `SoW_{CODE}_Detailed.md` (derived from `BDA PostAuction SoW.docx`, RFP Section 6)", "",
           "| # | Module | Type | Summary | Source |", "|---|--------|------|---------|--------|"]
    for m in MODULES:
        src = "Derived" if m["source"].startswith("Derived") else "Explicit"
        out.append(f"| {m['nn']} | {m['name']} | {m['type']} | {m['summary']} | {src} |")
    out += ["", "---", "", "## Module Details", ""]
    for m in MODULES:
        out += [f"### {m['nn']}. {m['name']}", "", f"- **Type:** {m['type']}", f"- **Summary:** {m['summary']}.",
                f"- **Source:** {m['source']}", f"- **In Scope:** {m['in_scope']}", f"- **Out of Scope:** {m['out_scope']}",
                f"- **Key Actors:** {m['actors']}", f"- **Dependencies:** {m['deps']}", f"- **Notes:** {m['notes'] or '—'}",
                "- **Features:** " + ", ".join(f"{f['id']} {f['name']}" for f in feats(m)), ""]
    out += ["## Excluded / Embedded Capabilities", "",
            "- **Workflow / approvals:** embedded in Modules 04 (result approval), 06 (extension, default, dispute), 07 (case processing) and 08 (deed approval); no separate workflow engine module.",
            "- **Document management:** embedded — site package (F03-004), letters and proceedings (F04-004, F05-003), payment proofs (F06-003), bidder uploads (F07-004), sale deed versions (F08-002), registered deed (F09-001).",
            "- **Scheduler / background jobs:** embedded in F03-001 (KPPP polling), F06-004 (daily reconciliation), F06-006 (reminders), F07-005 (timers), F12-002 (scheduled notifications).",
            "- **Live bidding, pre-auction, EMD and bidder registration:** out of scope (KPPP).",
            "- **VAPT execution and annual VAPT:** service by a CERT-In empanelled auditor; remediation is covered under warranty.",
            "- **Training, help guides and video tutorials:** service deliverables, outside feature hours.", "",
            "## Assumptions", "",
            "1. Access model is User → Role → Permissions with DS-wise data isolation; bidders use a separate OTP login.",
            "2. KPPP and CDRMS business connectors are built in the owning features (F03-001, F03-004, F09-004) on top of the Integration Hub framework (F13-001).",
            "3. The e-payment gateway connector and bank settlement intake are built in Module 06.",
            "4. Kaveri data exchange (F09-002) is conditional on permission from the Department of Stamps and Registration.",
            "5. BDA supplies approved templates, rules and legacy data in agreed formats.", "",
            "## Risks", "",
            "1. KPPP API availability and result format stability are outside BDA's and the vendor's control; the upload route mitigates this.",
            "2. CDRMS site package completeness affects case file quality and DS processing.",
            "3. Gateway settlement file formats vary by bank and may need rework.",
            "4. Tight timeline (development by T + 75 days) against the size of the scope.",
            "5. Legacy data quality for pending cases is unknown.", ""]
    (ROOT / f"{CODE}_Modules.md").write_text("\n".join(out))


def totals():
    all_h = sum(f["hours"] for f in FEATURES)
    go = sum(f["hours"] for f in FEATURES if f["notes"] == GO_LIVE)
    return all_h, go, all_h - go


def write_registry():
    all_h, go, opt = totals()
    reg = dict(project_code=CODE, project_name=PROJECT, client=CLIENT, output_root="Projects/BDA_Afterauction",
               source="BDA PostAuction SoW.docx", sow=f"SoW_{CODE}_Detailed.md",
               updated_at=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
               totals=dict(all_hours=all_h, go_live_hours=go, conditional_hours=opt, feature_count=len(FEATURES),
                           module_count=len(MODULES)),
               modules=[])
    for m in MODULES:
        reg["modules"].append(dict(nn=m["nn"], name=m["name"], slug=m["slug"], type=m["type"], folder=m["folder"],
                                   feature_count=len(feats(m)), total_hours=m["hours"],
                                   features=[dict(id=f["id"], seq=f["seq"], name=f["name"], slug=f["slug"], file=f["file"],
                                                  hours=f["hours"], notes=f["notes"], areas=f["areas"]) for f in feats(m)]))
    (ROOT / f"{CODE}_Registry.json").write_text(json.dumps(reg, indent=2, ensure_ascii=False) + "\n")


# ---------------------------------------------------------------- Step 6
def area_sum(fs):
    t = {a: 0 for a, _ in AREA_KEYS}
    for f in fs:
        for a in t:
            t[a] += f["areas"][a]
    return t


def pct(x, base):
    return f"{(100 * x / base):.1f}%" if base else "0.0%"


def write_summaries():
    all_h, go, opt = totals()
    by_type = {}
    for m in MODULES:
        by_type[m["type"]] = by_type.get(m["type"], 0) + m["hours"]
    areas = area_sum(FEATURES)
    cond = [f for f in FEATURES if f["notes"] != GO_LIVE]
    top = sorted(FEATURES, key=lambda f: -f["hours"])[:10]
    dev_heads = 4  # technical lead + 2 developers + QA, per RFP minimum team
    weeks_min = go / (dev_heads * 40)

    s = [f"# {CODE} — Estimation Summary (Internal)", "", f"**Project:** {PROJECT}  ", f"**Client:** {CLIENT}  ",
         f"**Project code:** {CODE}  ", f"**As of:** {TODAY}  ",
         f"**Source:** `SoW_{CODE}_Detailed.md` · feature indexes under `modules/*/features/_index.md`", "",
         "## Totals", "", "| Metric | Value |", "|--------|------:|", f"| Modules | {len(MODULES)} |", f"| Features | {len(FEATURES)} |",
         f"| **Recommended go-live hours** | **{go}** |", f"| Conditional hours | {opt} |", f"| **All-in hours** | **{all_h}** |",
         f"| All-in person-days (8 h) | {all_h / 8:.0f} |", "",
         "Conditional: " + ", ".join(f"{f['id']} {f['name']} ({f['hours']} h)" for f in cond) + ".", "",
         "## Per-module hours", "", "| # | Module | Type | Features | Hours | % of all-in |", "|---|--------|------|---------:|------:|------------:|"]
    for m in MODULES:
        s.append(f"| {m['nn']} | {m['name']} | {m['type']} | {len(feats(m))} | {m['hours']} | {pct(m['hours'], all_h)} |")
    s += ["", "## Hours by module type", "", "| Type | Hours | % |", "|------|------:|---:|"]
    s += [f"| {t} | {h} | {pct(h, all_h)} |" for t, h in sorted(by_type.items())]
    s += ["", "## Hours by estimation area", "", "| Area | Hours | % |", "|------|------:|---:|"]
    s += [f"| {a} | {h} | {pct(h, all_h)} |" for a, h in areas.items()]
    s += ["", "## Top 10 features by effort", "", "| Feature | Name | Hours |", "|---------|------|------:|"]
    s += [f"| {f['id']} | {f['name']} | {f['hours']} |" for f in top]
    s += ["", "## Timeline check against the RFP", "",
          f"The RFP expects development complete by T + 75 days and go-live by T + 100 days, with a minimum team of a project manager, "
          f"technical lead, two developers, a UI/UX designer and a QA engineer. Go-live effort of **{go} hours** equals about "
          f"**{weeks_min:.0f} weeks** for four build-and-test engineers at 40 hours per week. Meeting T + 75 days (about 11 weeks, "
          f"less SRS and prototype time) needs roughly **{go / (8 * 40):.0f} full-time engineers** over 8 build weeks. "
          "Plan a larger build team than the RFP minimum, reuse a proven workflow, notification and integration base, "
          "or agree a phased go-live (payments and tracking first, report builder and Kaveri later).", "",
          "## Not included in feature hours", "",
          "- Project management, SRS, integration study and UI/UX prototype phase",
          "- Training (two rounds per role), bidder help guide and video tutorials",
          "- VAPT by a CERT-In empanelled auditor (before go-live and annual)",
          "- On-site deployment of two professionals during implementation and support",
          "- Warranty (1 year) and operation and maintenance (3 years), including on-site support engineer",
          "- Hosting, gateway, SMS / WhatsApp / e-mail charges, DSC tokens and certificates", "",
          "## Open assumptions / risks", "",
          "1. KPPP API may not be provided; the controlled upload route (F03-002) is fully estimated as a fallback and both routes are in scope.",
          "2. CDRMS read API and status update API are available; otherwise a file-based exchange is needed.",
          "3. Payment gateway is BDA's existing provider with standard settlement files; one bank format assumed.",
          "4. e-Sign / DSC: server-side signing service or officer USB token signing through a signer utility, confirmed in SRS.",
          "5. Kaveri exchange (F09-002) is conditional and kept outside go-live hours.",
          "6. Legacy pending cases provided by BDA in the migration templates; data cleansing is BDA's responsibility.",
          "7. WhatsApp is built behind a switch and enabled only when BDA provides the business account.", "",
          "## Suggested next actions", "",
          "1. Review modules, features and hours with the delivery lead.",
          "2. Decide team size and phasing against the T + 100 days go-live (see timeline check).",
          "3. Confirm KPPP API, CDRMS API, gateway, e-Sign and Kaveri positions during the integration study.",
          f"4. Share the client Word export: `{CODE}_Estimation_Summary_FeatureWise_Client.docx`.", ""]
    (ROOT / f"{CODE}_Estimation_Summary.md").write_text("\n".join(s))

    fw = [f"# {CODE} — Feature-wise Estimation Summary (Internal)", "", f"**Project:** {PROJECT}  ", f"**Project code:** {CODE}  ",
          f"**As of:** {TODAY}", "", "## Grand totals", "", "| Bucket | Hours |", "|--------|------:|",
          f"| Recommended go-live | {go} |", f"| Conditional | {opt} |", f"| All-in | {all_h} |",
          f"| Features | {len(FEATURES)} |", f"| Modules | {len(MODULES)} |", "",
          "## Module-wise summary", "", "| # | Module | Type | Features | UI | API | DB | Integ. | Test | Docs | Hours |",
          "|---|--------|------|---------:|---:|---:|---:|---:|---:|---:|------:|"]
    for m in MODULES:
        a = area_sum(feats(m))
        fw.append(f"| {m['nn']} | {m['name']} | {m['type']} | {len(feats(m))} | {a['UI']} | {a['API']} | {a['Database']} | "
                  f"{a['Integration']} | {a['Testing']} | {a['Documentation']} | {m['hours']} |")
    fw.append(f"| | **Total** | | **{len(FEATURES)}** | **{areas['UI']}** | **{areas['API']}** | **{areas['Database']}** | "
              f"**{areas['Integration']}** | **{areas['Testing']}** | **{areas['Documentation']}** | **{all_h}** |")
    fw += ["", "## Feature-wise detail", ""]
    for m in MODULES:
        fw += [f"### {m['nn']} — {m['name']}", "",
               "| Feature ID | Feature Name | Screens | APIs | UI | API | DB | Integ. | Test | Docs | Hours | Notes |",
               "|------------|--------------|--------:|-----:|---:|---:|---:|---:|---:|---:|------:|-------|"]
        for f in feats(m):
            a = f["areas"]
            fw.append(f"| {f['id']} | {f['name']} | {len(f['screens'])} | {len(f['apis'])} | {a['UI']} | {a['API']} | {a['Database']} | "
                      f"{a['Integration']} | {a['Testing']} | {a['Documentation']} | {f['hours']} | {f['notes']} |")
        fw += ["", f"**Module total: {m['hours']} hours**", ""]
    fw += ["## Effort profile", "", "| Area | Hours | % |", "|------|------:|---:|"]
    fw += [f"| {a} | {h} | {pct(h, all_h)} |" for a, h in areas.items()]
    fw += ["", "## Indicative phasing (from the SOW)", "",
           "| Phase | Timeline | Modules |", "|-------|----------|---------|",
           "| SRS, integration study, prototype | T + 21 to T + 30 days | All (design) |",
           "| Foundation | T + 30 to T + 45 days | 01, 02, 12, 13, 14, 15 |",
           "| Intake, approval, bidder account | T + 40 to T + 60 days | 03, 04, 05 |",
           "| Payments and DS processing | T + 50 to T + 75 days | 06, 07, 08, 09 |",
           "| Portal, dashboards, migration | T + 60 to T + 90 days | 10, 11, 16 |",
           "| UAT, VAPT, go-live | T + 75 to T + 100 days | All |", "",
           "## Exclusions", "",
           "- PM, SRS, prototype, training, help material, VAPT service, on-site staffing, warranty and O&M, hosting and gateway charges", ""]
    (ROOT / f"{CODE}_Estimation_Summary_FeatureWise.md").write_text("\n".join(fw))


# ---------------------------------------------------------------- DOCX helpers
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

NAVY = RGBColor(0x1B, 0x3A, 0x5F)


def base_doc():
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = NAVY
        doc.styles[name].font.name = "Calibri"
    return doc


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_runs(par, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = par.add_run(part.replace("`", ""))
        run.bold = i % 2 == 1


def add_table(doc, header, body, num_cols=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        shade(c, "1B3A5F")
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_runs(p, str(v))
            for r in p.runs:
                r.font.size = Pt(9.5)
            if i in num_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if ri % 2 == 1:
                shade(cells[i], "F2F4F7")
    doc.add_paragraph()
    return t


def md_to_docx(md_path: Path, out_path: Path):
    doc = base_doc()
    lines = md_path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            body = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                body.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, header, body)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2 if re.match(r"^\d{2}\. ", line[4:]) else 3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(20)
            r.font.color.rgb = NAVY
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip() in ("---", ""):
            pass
        else:
            add_runs(doc.add_paragraph(), line.rstrip(" "))
        i += 1
    doc.save(out_path)


# ---------------------------------------------------------------- Step 7
CLIENT_NOTES = {}


def write_client_docx():
    all_h, go, opt = totals()
    doc = base_doc()
    for text, size, bold in [("Module & Feature-wise Effort Estimate", 22, True), (PROJECT, 14, True),
                             (f"Prepared for {CLIENT} · Client Copy · Effort in person-hours", 11, False),
                             (f"Date: {TODAY}", 10, False)]:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("This document sets out the module and feature-wise effort for the BDA Auction Module (Post-Auction). The Module "
                      "takes each awarded site from the KPPP auction result to the completion of the auction process: result intake and "
                      "verification, bidder intimation, instalment collection and reconciliation, Deputy Secretary processing, draft sale "
                      "deed, registration recording and completion, with a bilingual Bidder Tracking Portal, dashboards and reports.")
    doc.add_paragraph("Hours cover the design and build of the application software, its integrations, security and audit controls, "
                      "operational tooling and migration utilities, with associated testing and technical documentation.")
    doc.add_paragraph("Hours do not cover project management, the requirement study and prototype phase, training and help material, the "
                      "VAPT service, on-site staffing, warranty and operation and maintenance, or hosting and third-party service charges.")

    doc.add_heading("2. Summary Totals", level=1)
    add_table(doc, ["Bucket", "Person-hours"], [["Recommended for go-live", go], ["Conditional (subject to departmental permission)", opt],
                                                ["**Total (all items)**", f"**{all_h}**"], ["Modules", len(MODULES)],
                                                ["Features", len(FEATURES)]], num_cols=(1,))

    doc.add_heading("3. Module-wise Summary", level=1)
    add_table(doc, ["Module", "Type", "Features", "Hours", "% of go-live"],
              [[f"{m['nn']}. {m['name']}", m["type"], len(feats(m)), m["hours"],
                pct(sum(f["hours"] for f in feats(m) if f["notes"] == GO_LIVE), go)] for m in MODULES]
              + [["**Total**", "", f"**{len(FEATURES)}**", f"**{all_h}**", "**100%**"]], num_cols=(2, 3, 4))

    doc.add_heading("4. Feature-wise Detail", level=1)
    for m in MODULES:
        doc.add_heading(f"{m['nn']}. {m['name']}", level=2)
        doc.add_paragraph(m["summary"] + ".")
        add_table(doc, ["Feature ID", "Feature Name", "Hours", "Notes"],
                  [[f["id"], f["name"], f["hours"], f["client"]] for f in feats(m)]
                  + [["", "**Module total**", f"**{m['hours']}**", ""]], num_cols=(2,))

    doc.add_heading("5. Conditional Items and Commercial Line Items", level=1)
    add_table(doc, ["Item", "Basis"], [
        ["F09-002 Kaveri Registration Data Exchange", f"{opt} person-hours, taken up only if the Department of Stamps and Registration permits data exchange"],
        ["Operation and maintenance for three years from go-live", "Quoted separately as part (b) of the financial proposal"],
        ["On-site team during implementation and support", "As per the minimum team in the RFP"],
        ["Change of scope", "At the man-month rate quoted in the financial proposal"],
    ])

    doc.add_heading("6. Effort Profile", level=1)
    areas = area_sum(FEATURES)
    label = {"UI": "User interface (screens and portal pages)", "API": "Application services and business rules",
             "Database": "Data model and storage", "Integration": "Integrations", "Testing": "Testing", "Documentation": "Documentation"}
    add_table(doc, ["Work area", "Hours", "% of total"], [[label[a], h, pct(h, all_h)] for a, h in areas.items()], num_cols=(1, 2))
    top = sorted(FEATURES, key=lambda f: -f["hours"])[:10]
    doc.add_paragraph("Largest features by effort:")
    add_table(doc, ["Feature ID", "Feature Name", "Hours"], [[f["id"], f["name"], f["hours"]] for f in top], num_cols=(2,))
    by_type = {}
    for m in MODULES:
        by_type[m["type"]] = by_type.get(m["type"], 0) + m["hours"]
    add_table(doc, ["Module type", "Hours", "% of total"], [[t, h, pct(h, all_h)] for t, h in sorted(by_type.items())], num_cols=(1, 2))

    doc.add_heading("7. Indicative Delivery Phasing", level=1)
    add_table(doc, ["Phase", "Timeline", "Scope"], [
        ["Requirements and design", "T + 21 to T + 30 days", "System Requirement Specification, KPPP and CDRMS integration study, payment and area rules, UI/UX prototype"],
        ["Foundation", "T + 30 to T + 45 days", "Access and roles, rules and templates, notifications, integration services, audit and security, environments"],
        ["Result to intimation", "T + 40 to T + 60 days", "KPPP intake and site package, verification and approval, bidder account and intimation"],
        ["Payments to completion", "T + 50 to T + 75 days", "Instalments and reconciliation, DS processing, draft sale deed, registration and completion"],
        ["Portal, MIS and migration", "T + 60 to T + 90 days", "Bidder Tracking Portal, dashboards and reports, legacy case migration"],
        ["Acceptance and go-live", "T + 75 to T + 100 days", "UAT, VAPT clearance, training and go-live; final acceptance at T + 160 days"],
    ])

    doc.add_heading("8. Items Not Included in Feature Hours", level=1)
    for t in ["Project management, System Requirement Specification and UI/UX prototype",
              "Training for BDA officials (minimum two rounds), bidder help guide and video tutorials",
              "VAPT by a CERT-In empanelled auditor before go-live and annually",
              "On-site deployment of professionals during implementation and support",
              "Warranty (one year) and operation and maintenance (three years)",
              "Hosting infrastructure, payment gateway, SMS, WhatsApp and e-mail charges, DSC tokens and certificates",
              "Live bidding, pre-auction activities, EMD handling and bidder registration (performed on KPPP)",
              "Registration of the sale deed (Department of Stamps and Registration)"]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("9. Key Planning Assumptions", level=1)
    for t in ["KPPP provides auction results through an API; where it does not, standard KPPP result reports are uploaded through the controlled upload facility, which is included.",
              "CDRMS provides read access to the site package and accepts completion status updates.",
              "BDA provides the e-payment gateway, bank settlement data, e-mail, SMS (with DLT registration) and WhatsApp accounts, and the e-Sign / DSC arrangement.",
              "Approving officers hold valid DSC tokens.",
              "BDA provides approved templates for letters, proceedings, notices and the draft sale deed, and the payment and area allocation rules.",
              "Legacy pending cases are provided by BDA in the agreed migration templates.",
              "Hosting is provided at the Karnataka State Data Centre or on BDA-specified infrastructure.",
              "Access is role-based with Deputy Secretary-wise data isolation; bidders use a separate OTP login."]:
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("10. Disclaimer", level=1)
    doc.add_paragraph("Estimates are indicative for planning and commercial discussion. Final effort may vary with the KPPP and CDRMS "
                      "integration positions confirmed in the requirement study, the payment gateway and bank formats, the e-Sign approach, "
                      "the volume and quality of legacy data, and any scope changes agreed with BDA.")
    doc.save(ROOT / f"{CODE}_Estimation_Summary_FeatureWise_Client.docx")


def main():
    validate()
    write_features()
    write_modules_md()
    write_registry()
    write_summaries()
    md_to_docx(ROOT / f"SoW_{CODE}_Detailed.md", ROOT / f"SoW_{CODE}_Detailed.docx")
    write_client_docx()
    all_h, go, opt = totals()
    print(f"Features {len(FEATURES)} · go-live {go} · conditional {opt} · all-in {all_h}")
    for m in MODULES:
        print(f"  {m['nn']} {m['name']}: {len(feats(m))} features, {m['hours']} h")


if __name__ == "__main__":
    main()
