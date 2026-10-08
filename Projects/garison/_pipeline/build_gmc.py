"""Build GMC estimation artefacts (Steps 2-3, 5, 6, 7) from feature data.

Run from repo root:
    .venv/bin/python Projects/garison/_pipeline/build_gmc.py
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

from gmc_data_a import FEATURES as FA  # noqa: E402
from gmc_data_b import FEATURES as FB  # noqa: E402
from gmc_data_c import FEATURES as FC  # noqa: E402

FEATURES = FA + FB + FC
CODE = "GMC"
PROJECT = "Garrison Online Moulding Catalogue Platform"
CLIENT = "Garrison Industries"
TODAY = dt.date.today().isoformat()

CX = {"S": "Simple", "E": "Easy", "M": "Medium", "C": "Complex", "V": "Very High"}
HOURS = {
    "UI": {"S": 2, "E": 4, "M": 6, "C": 12, "V": 16},
    "API": {"S": 2, "E": 4, "M": 6, "C": 8, "V": 16},
    "Database": {"S": 1, "E": 2, "M": 4, "C": 6, "V": 10},
    "Integration": {"S": 4, "E": 8, "M": 16, "C": 24, "V": 40},
    "Testing": {"S": 2, "E": 4, "M": 6, "C": 8, "V": 12},
    "Documentation": {"S": 1, "E": 2, "M": 3, "C": 4, "V": 6},
}
AREA_KEYS = [("UI", "screens"), ("API", "apis"), ("Database", "db"), ("Integration", "integ"), ("Testing", "tests"), ("Documentation", "docs")]

MODULES = [
    dict(nn="01", name="Administrator Access & Role Management", type="Shared",
         summary="Admin sign-in and session security, user management, configurable role permissions, audit log",
         source="Explicit (Proposal §20, §22; FR-14, FR-20)",
         in_scope="Admin authentication, password policy, lockout, reset; user invite and lifecycle; role-permission matrix; audit log with field-level diff",
         out_scope="Public visitor accounts; SSO; MFA (future option)",
         actors="Super Administrator; Catalogue Administrator; Catalogue Reviewer; Content Administrator",
         deps="F06-005 (email connector); F13-002 (rate limiting)",
         notes="Access model User → Role → Permissions; four default roles seeded"),
    dict(nn="02", name="Catalogue Taxonomy & Master Data", type="Business",
         summary="Categories and sub-categories, styles, collections, materials/species, finishes",
         source="Explicit (Proposal §8; FR-02, FR-07, FR-08)",
         in_scope="Category hierarchy with ordering, imagery and SEO intro; styles; curated collections; material and finish masters with merge",
         out_scope="Multi-language taxonomy",
         actors="Catalogue Administrator; Content Administrator",
         deps="Module 03 (templates); F11-001 (redirects); F09-001 (re-index)",
         notes="Classification layers kept independent for accurate filtering"),
    dict(nn="03", name="Dynamic Attribute & Dimension Framework", type="Business",
         summary="Attribute library, category attribute templates, template change impact, dimension normalisation",
         source="Explicit (Proposal §7; FR-04, FR-05)",
         in_scope="Typed attributes with units and permitted values; per-category templates with mandatory/conditional/filterable flags; impact management; fraction/decimal/metric parsing",
         out_scope="Formula-based calculated attributes",
         actors="Catalogue Administrator",
         deps="Module 02; consumed by Modules 04, 06, 09, 12",
         notes="Core enabler of 'no code change for new categories'"),
    dict(nn="04", name="Product Profile Management", type="Business",
         summary="Product register, core profile, dynamic dimensions, materials/finishes/lengths, media, documents, lifecycle, preview",
         source="Explicit (Proposal §6, §12.1–12.2; FR-03, FR-05, FR-06)",
         in_scope="Product workspace; template-driven entry; typed media and drawings; technical PDFs; duplicate/archive/restore; draft preview",
         out_scope="Pricing, stock, CAD downloads",
         actors="Catalogue Administrator; Content Administrator; Catalogue Reviewer (preview)",
         deps="Modules 02, 03, 06; F13-001",
         notes="Working-draft model keeps published pages live during edits"),
    dict(nn="05", name="Combinations & Product Relationships", type="Business",
         summary="Typed product relationships; combinations as separate entity with component assembly",
         source="Explicit (Proposal §11; FR-09, FR-10)",
         in_scope="Related/matching/alternative/similar relationships with reciprocity; combination records; component roles and order; 'Used in combinations'",
         out_scope="Automatic recommendations",
         actors="Catalogue Administrator",
         deps="Module 04; Module 06; Module 08",
         notes="Combinations follow the same workflow as products"),
    dict(nn="06", name="Content Governance & Publishing Workflow", type="Business",
         summary="Completeness and duplicate rules, submit/review/approve, publish/unpublish, workflow emails",
         source="Explicit (Proposal §12.3–12.4; FR-15, FR-21)",
         in_scope="Completeness score and gate; duplicate detection; review queue; maker-checker decisions; bulk publish; email connector and templates",
         out_scope="Multi-level approval chains; scheduled publishing",
         actors="Catalogue Administrator; Content Administrator; Catalogue Reviewer; Super Administrator",
         deps="Modules 04, 05, 07; F09-001; F11-002",
         notes="Notification embedded here (F06-005); no separate notification module"),
    dict(nn="07", name="External Source, Provenance & Content Rights", type="Business",
         summary="External source register, per-product provenance, rights-based publication gate and dashboard",
         source="Explicit (Proposal §14; FR-16)",
         in_scope="Source master with agreements; provenance fields and status history; gate matrix; Garrison-owned-content-only mode; re-verification",
         out_scope="Automated scraping of external catalogues; legal verification of rights",
         actors="Catalogue Administrator; Super Administrator; Catalogue Reviewer (view)",
         deps="Modules 04, 06",
         notes="Gate enforced server-side and cannot be bypassed"),
    dict(nn="08", name="Public Catalogue Experience", type="Business",
         summary="Catalogue home, listings, product detail, combination detail, responsive design system",
         source="Explicit (Proposal §5, §10, §15; FR-01, FR-13)",
         in_scope="Home; category/style/collection listings; product page with drawing viewer; combination page; brand-aligned component library; accessibility baseline",
         out_scope="Customer accounts, favourites, comparison",
         actors="Visitor; Content Administrator (home settings)",
         deps="Modules 02–05, 09, 10, 11",
         notes="Server-rendered/pre-rendered for SEO"),
    dict(nn="09", name="Search & Filtering", type="Business",
         summary="Search index sync, ranked keyword/code/dimension search, configurable facets, suggestions",
         source="Explicit (Proposal §9; FR-11, FR-12)",
         in_scope="Dedicated search index; code-priority ranking; dimension phrase parsing; synonyms; template-driven list and range facets; type-ahead",
         out_scope="AI/visual search",
         actors="Visitor; Catalogue Administrator (synonyms, facets)",
         deps="Modules 03, 06; F13-003",
         notes="Search engine choice confirmed at discovery"),
    dict(nn="10", name="Product Enquiry", type="Business",
         summary="Contextual enquiry form, routing and acknowledgement emails, basic enquiry register",
         source="Explicit (Proposal FR-17)",
         in_scope="Enquiry form with spam protection; routing rules; acknowledgement; register with status, export and retention",
         out_scope="Full enquiry/quote management; CRM integration",
         actors="Visitor; Catalogue Administrator; Garrison recipients",
         deps="Module 08; F06-005",
         notes="Full enquiry management is a future enhancement"),
    dict(nn="11", name="SEO, Analytics & Website Integration", type="Shared",
         summary="SEO URLs/metadata/redirects, sitemap and structured data, Garrison website integration, usage analytics",
         source="Explicit (Proposal §21; FR-18, FR-19)",
         in_scope="URL scheme, metadata, canonical, redirects; sitemap, JSON-LD, social tags; header/footer/navigation integration; analytics events and insight dashboard",
         out_scope="Ongoing SEO content marketing",
         actors="Content Administrator; Catalogue Administrator; Super Administrator; Garrison web administrator",
         deps="Module 08; Garrison website and hosting access",
         notes="Integration method (sub-path/sub-domain) confirmed at discovery"),
    dict(nn="12", name="Catalogue Data Import & Bulk Operations", type="Technical",
         summary="Import templates and validation, bulk media matching, staging and commit, optional export/bulk update",
         source="Explicit (Proposal §13; FR-22); F12-004 Derived",
         in_scope="Category import templates; validated bulk import; media auto-matching; staging grid and commit/rollback; export and bulk update (optional)",
         out_scope="Manual data extraction and entry labour (priced per product)",
         actors="Catalogue Administrator; Data population team",
         deps="Modules 03, 04, 06; F13-001, F13-002",
         notes="Tooling for the population workstream; labour is a separate commercial line"),
    dict(nn="13", name="Platform Security, Performance & Operations", type="Technical",
         summary="Media processing and CDN, security hardening, backup and monitoring, environments and deployment",
         source="Derived (Proposal §16, §19, §22, §26)",
         in_scope="Derivatives and CDN; OWASP controls and upload scanning; backups, monitoring, logging; dev/staging/prod and pipeline",
         out_scope="Hosting charges; independent VAPT; post-warranty operations (AMC)",
         actors="Super Administrator; Delivery/support team",
         deps="Garrison hosting environment",
         notes="Hosting stack finalised after discovery"),
]


def slugify(text: str) -> str:
    text = text.lower().replace("&", "and")
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


for m in MODULES:
    m["slug"] = slugify(m["name"])
    m["folder"] = f"modules/{m['nn']}-{m['slug']}"


def rows(f):
    out = []
    for area, key in AREA_KEYS:
        for name, cx in f.get(key, []):
            out.append((area, name, CX[cx], HOURS[area][cx]))
    return out


def area_totals(f):
    t = {a: 0 for a, _ in AREA_KEYS}
    for area, _, _, h in rows(f):
        t[area] += h
    return t


for f in FEATURES:
    f.setdefault("notes", "Go-live")
    f["slug"] = slugify(f["name"])
    f["nn"] = f["id"][1:3]
    f["seq"] = f["id"][4:]
    f["hours"] = sum(r[3] for r in rows(f))
    f["areas"] = area_totals(f)


def module_of(f):
    return next(m for m in MODULES if m["nn"] == f["nn"])


def feats(m):
    return [f for f in FEATURES if f["nn"] == m["nn"]]


def bullets(items):
    return "\n".join(f"- {i}" for i in items) if items else "- None"


# ---------------------------------------------------------------- validation
def validate():
    errors = []
    sow = (ROOT / f"SoW_{CODE}_Detailed.md").read_text()
    sow_feats = dict(re.findall(r"^\| (F\d{2}-\d{3}) \| ([^|]+?) \|", sow, re.M))
    ids = [f["id"] for f in FEATURES]
    if len(ids) != len(set(ids)):
        errors.append("duplicate feature IDs")
    for f in FEATURES:
        if f["id"] not in sow_feats:
            errors.append(f"{f['id']} missing in SOW")
        elif sow_feats[f["id"]].strip() != f["name"]:
            errors.append(f"{f['id']} name mismatch: SOW '{sow_feats[f['id']].strip()}' vs '{f['name']}'")
        if len(f["journey"]) != 8:
            errors.append(f"{f['id']} journey has {len(f['journey'])} points")
        for k in ("desc", "obj", "actors", "caps", "inputs", "outputs", "rules", "valid", "tests", "docs"):
            if not f.get(k):
                errors.append(f"{f['id']} missing {k}")
    for fid in sow_feats:
        if fid not in ids:
            errors.append(f"{fid} in SOW but no feature data")
    for m in MODULES:
        if f"### {m['nn']}. {m['name']}" not in sow:
            errors.append(f"module heading missing in SOW: {m['nn']}. {m['name']}")
    if errors:
        raise SystemExit("Validation failed:\n" + "\n".join(errors))


# ---------------------------------------------------------------- Step 5
def render_feature(f):
    m = module_of(f)
    lines = [
        f"# {f['id']} — {f['name']}",
        "",
        "## Feature Specification",
        "",
        "| Field | Value |",
        "|-------|-------|",
        f"| Feature ID | {f['id']} |",
        f"| Feature Name | {f['name']} |",
        f"| Module | {m['nn']} — {m['name']} |",
        f"| Module Type | {m['type']} |",
        f"| Release | {f['notes']} |",
        "",
        "### Description",
        f["desc"],
        "",
        "### Business Objective",
        f["obj"],
        "",
        "### Actors",
        bullets(f["actors"]),
        "",
        "### User Journey",
    ]
    labels = ["Who creates", "How they create", "How they edit", "How they send for approval",
              "How the approver sees pending items", "How the approver decides",
              "How the maker sees history", "How the approver / auditor sees history"]
    for i, (lab, txt) in enumerate(zip(labels, f["journey"]), 1):
        lines.append(f"{i}. **{lab}:** {txt}")
    lines += ["", "### Functional Capabilities", bullets(f["caps"]), ""]
    lines += ["### UI Screens", bullets([s for s, _ in f["screens"]]) if f["screens"] else "- None (no user interface artefacts in this feature)", ""]
    lines += ["### Logical API Operations", bullets([a for a, _ in f["apis"]]) if f["apis"] else "- None (no application operations in this feature)", ""]
    sections = [("Inputs", "inputs"), ("Outputs", "outputs"), ("Business Rules", "rules"), ("Validations", "valid"),
                ("Calculations", "calc")]
    for title, key in sections:
        lines += [f"### {title}", bullets(f[key]), ""]
    lines += ["### Workflow", f["flow"], ""]
    for title, key in [("Search Requirements", "search"), ("Documents", "documents"), ("Reports", "reports"),
                       ("Notifications", "notif"), ("Security", "sec"), ("Integrations", "ints"),
                       ("Dependencies", "deps"), ("Assumptions", "assume"), ("Risks", "risks"),
                       ("Non Functional Requirements", "nfr")]:
        lines += [f"### {title}", bullets(f[key]), ""]
    lines += ["## Estimation Breakdown", "", "| Area | Artefact | Complexity | Hours |", "|------|----------|------------|------:|"]
    for area, name, cx, h in rows(f):
        lines.append(f"| {area} | {name} | {cx} | {h} |")
    lines += ["", f"**Feature Total: {f['hours']} hours**", ""]
    a = f["areas"]
    lines += ["| UI | API | Database | Integration | Testing | Documentation | Total |",
              "|---:|---:|---:|---:|---:|---:|---:|",
              f"| {a['UI']} | {a['API']} | {a['Database']} | {a['Integration']} | {a['Testing']} | {a['Documentation']} | **{f['hours']}** |",
              "", "### Estimation Assumptions",
              "- Every UI screen and logical API operation listed above is estimated as its own row using the Estimation Master.",
              "- Enterprise-grade implementation: validation, error handling, logging, security checks and unit tests included in each row."]
    lines += [f"- {e}" for e in f["est"]]
    if not f["integ"]:
        lines.append("- No external integration channel in this feature; Integration row omitted.")
    return "\n".join(lines) + "\n"


def write_features():
    for m in MODULES:
        fdir = ROOT / m["folder"] / "features"
        fdir.mkdir(parents=True, exist_ok=True)
        for old in fdir.glob("F*.md"):
            old.unlink()
        ms = feats(m)
        for f in ms:
            f["file"] = f"{m['folder']}/features/{f['id']}-{f['slug']}.md"
            (ROOT / f["file"]).write_text(render_feature(f))
        m["hours"] = sum(f["hours"] for f in ms)
        m["golive"] = sum(f["hours"] for f in ms if f["notes"] == "Go-live")
        idx = [f"# Feature Index", "", f"## {m['nn']} — {m['name']}", "",
               "| Feature ID | Feature Name | Total Hours | Notes |", "|------------|--------------|------------:|-------|"]
        idx += [f"| {f['id']} | {f['name']} | {f['hours']} | {f['notes']} |" for f in ms]
        idx += ["", f"**Module Total:** **{m['hours']} hours**", ""]
        (fdir / "_index.md").write_text("\n".join(idx))
        readme = [f"# {m['nn']}. {m['name']}", "", f"Type: {m['type']}", f"Summary: {m['summary']}.", "",
                  f"- **Source:** {m['source']}", f"- **In Scope:** {m['in_scope']}", f"- **Out of Scope:** {m['out_scope']}",
                  f"- **Key Actors:** {m['actors']}", f"- **Dependencies:** {m['deps']}", f"- **Notes:** {m['notes']}", "",
                  f"Features: {len(ms)} · **{m['hours']} hours** (see `features/_index.md`)", ""]
        (ROOT / m["folder"] / "README.md").write_text("\n".join(readme))


# ---------------------------------------------------------------- Steps 2-3
def write_modules_md():
    out = [f"# {PROJECT} — Module List", "", f"Source: `SoW_{CODE}_Detailed.md` (derived from `Garrison_Technical_Proposal.md`)", "",
           "| # | Module | Type | Summary | Source |", "|---|--------|------|---------|--------|"]
    for m in MODULES:
        src = "Derived" if m["source"].startswith("Derived") else "Explicit"
        out.append(f"| {m['nn']} | {m['name']} | {m['type']} | {m['summary']} | {src} |")
    out += ["", "---", "", "## Module Details", ""]
    for m in MODULES:
        out += [f"### {m['nn']}. {m['name']}", "", f"- **Type:** {m['type']}", f"- **Summary:** {m['summary']}.",
                f"- **Source:** {m['source']}", f"- **In Scope:** {m['in_scope']}", f"- **Out of Scope:** {m['out_scope']}",
                f"- **Key Actors:** {m['actors']}", f"- **Dependencies:** {m['deps']}", f"- **Notes:** {m['notes']}",
                f"- **Features:** " + ", ".join(f"{f['id']} {f['name']}" for f in feats(m)), ""]
    out += ["## Embedded capabilities (no separate module)", "",
            "- **Notification:** workflow emails in F06-005 (shared email connector); enquiry emails in F10-002.",
            "- **Document management:** technical documents in F04-005; source agreements in F07-001.",
            "- **Workflow:** catalogue review and publishing workflow in Module 06.",
            "- **Audit:** cross-cutting audit log in F01-004.", ""]
    (ROOT / f"{CODE}_Modules.md").write_text("\n".join(out))


def totals():
    all_h = sum(f["hours"] for f in FEATURES)
    go = sum(f["hours"] for f in FEATURES if f["notes"] == "Go-live")
    return all_h, go, all_h - go


def write_registry():
    all_h, go, opt = totals()
    reg = dict(project_code=CODE, project_name=PROJECT, client=CLIENT, output_root="Projects/garison",
               source="Garrison_Technical_Proposal.md", sow=f"SoW_{CODE}_Detailed.md",
               updated_at=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
               totals=dict(all_hours=all_h, go_live_hours=go, optional_hours=opt, feature_count=len(FEATURES), module_count=len(MODULES)),
               modules=[])
    for m in MODULES:
        reg["modules"].append(dict(nn=m["nn"], name=m["name"], slug=m["slug"], type=m["type"], folder=m["folder"],
                                   feature_count=len(feats(m)), total_hours=m["hours"],
                                   features=[dict(id=f["id"], seq=f["seq"], name=f["name"], slug=f["slug"], file=f["file"],
                                                  hours=f["hours"], notes=f["notes"], areas=f["areas"]) for f in feats(m)]))
    (ROOT / f"{CODE}_Registry.json").write_text(json.dumps(reg, indent=2, ensure_ascii=False) + "\n")


# ---------------------------------------------------------------- Step 6
def area_sum(fs):
    t = {a: 0 for a, _ in AREA_KEYS}
    for f in fs:
        for a in t:
            t[a] += f["areas"][a]
    return t


def pct(x, base):
    return f"{(100 * x / base):.1f}%" if base else "0.0%"


def write_summaries():
    all_h, go, opt = totals()
    by_type = {}
    for m in MODULES:
        by_type[m["type"]] = by_type.get(m["type"], 0) + m["hours"]
    areas = area_sum(FEATURES)
    optional = [f for f in FEATURES if f["notes"] != "Go-live"]

    s = [f"# {CODE} — Estimation Summary (Internal)", "", f"**Project:** {PROJECT}  ", f"**Client:** {CLIENT}  ",
         f"**Project code:** {CODE}  ", f"**As of:** {TODAY}  ",
         f"**Source:** `SoW_{CODE}_Detailed.md` · feature indexes under `modules/*/features/_index.md`", "",
         "## Totals", "", "| Metric | Value |", "|--------|------:|", f"| Modules | {len(MODULES)} |", f"| Features | {len(FEATURES)} |",
         f"| **Recommended go-live hours** | **{go}** |", f"| Optional / conditional hours | {opt} |", f"| **All-in hours** | **{all_h}** |", "",
         "Optional hours: " + ", ".join(f"{f['id']} {f['name']} ({f['hours']} h)" for f in optional) + ".", "",
         "## Per-module hours", "", "| # | Module | Type | Features | Hours | % of all-in |", "|---|--------|------|---------:|------:|------------:|"]
    for m in MODULES:
        s.append(f"| {m['nn']} | {m['name']} | {m['type']} | {len(feats(m))} | {m['hours']} | {pct(m['hours'], all_h)} |")
    s += ["", "## Hours by module type", "", "| Type | Hours |", "|------|------:|"]
    s += [f"| {t} | {h} |" for t, h in sorted(by_type.items())]
    s += ["", "## Hours by estimation area", "", "| Area | Hours | % |", "|------|------:|---:|"]
    s += [f"| {a} | {h} | {pct(h, all_h)} |" for a, h in areas.items()]
    weeks_5 = go / (5 * 40)
    weeks_6 = go / (6 * 40)
    s += ["", "## Timeline check against the proposal", "",
          f"The proposal states an indicative 10–12 week duration. Go-live effort of **{go} hours** corresponds to about "
          f"**{weeks_5:.1f} weeks with a 5-person team** or **{weeks_6:.1f} weeks with a 6-person team** at 40 hours per person-week, "
          f"before project management, UX/UI design and population labour. Delivering in 12 weeks would need about **{go / (12 * 40):.1f} full-time engineers**; "
          "otherwise extend the timeline to roughly 17–20 weeks or trim scope (candidates: F12-004, F11-004 insight dashboard, F07-003 dashboard, "
          "F09-004 suggestions, F03-003 impact management).", "",
          "## Not included in feature hours", "",
          "- Project management, discovery workshops and UX/UI design phase (wireframes, visual design)",
          "- Initial catalogue population labour (extraction, normalisation, drawing preparation) — commercial lines B and C, per product",
          "- Hosting, CDN, email, analytics and search service subscriptions",
          "- Independent penetration testing; AMC / post-warranty support (commercial line D)", "",
          "## Open assumptions / risks", "",
          "1. Access model User → Role → Permissions with four default roles; no public login.",
          "2. Dedicated search engine available in the hosting environment (choice at discovery).",
          "3. Website integration via sub-path or sub-domain; CMS plugin embedding would need re-estimation.",
          "4. Garrison confirms third-party content rights; the platform records and enforces, but does not verify legally.",
          "5. Initial product count unknown; population labour priced separately.",
          "6. English only; inches primary with optional millimetre display.",
          "7. No e-commerce, pricing, inventory, CRM or ERP.", "",
          "## Suggested next actions", "",
          "1. Review modules, features and hours with the delivery lead; decide on F12-004 (optional).",
          "2. Reconcile hours with the proposal's 10–12 week timeline and team plan (see timeline check).",
          "3. Confirm website platform, hosting and search engine in discovery.",
          f"4. Use the client Word export: `{CODE}_Estimation_Summary_FeatureWise_Client.docx`.", ""]
    (ROOT / f"{CODE}_Estimation_Summary.md").write_text("\n".join(s))

    fw = [f"# {CODE} — Feature-wise Estimation Summary (Internal)", "", f"**Project:** {PROJECT}  ", f"**Project code:** {CODE}  ",
          f"**As of:** {TODAY}", "", "## Grand totals", "", "| Bucket | Hours |", "|--------|------:|",
          f"| Recommended go-live | {go} |", f"| Optional / conditional | {opt} |", f"| All-in | {all_h} |",
          f"| Features | {len(FEATURES)} |", f"| Modules | {len(MODULES)} |", "",
          "## Module-wise summary", "", "| # | Module | Type | Features | UI | API | DB | Integ. | Test | Docs | Hours |",
          "|---|--------|------|---------:|---:|---:|---:|---:|---:|---:|------:|"]
    for m in MODULES:
        a = area_sum(feats(m))
        fw.append(f"| {m['nn']} | {m['name']} | {m['type']} | {len(feats(m))} | {a['UI']} | {a['API']} | {a['Database']} | "
                  f"{a['Integration']} | {a['Testing']} | {a['Documentation']} | {m['hours']} |")
    fw.append(f"| | **Total** | | **{len(FEATURES)}** | **{areas['UI']}** | **{areas['API']}** | **{areas['Database']}** | "
              f"**{areas['Integration']}** | **{areas['Testing']}** | **{areas['Documentation']}** | **{all_h}** |")
    fw += ["", "## Feature-wise detail", ""]
    for m in MODULES:
        fw += [f"### {m['nn']} — {m['name']}", "",
               "| Feature ID | Feature Name | Screens | APIs | UI | API | DB | Integ. | Test | Docs | Hours | Notes |",
               "|------------|--------------|--------:|-----:|---:|---:|---:|---:|---:|---:|------:|-------|"]
        for f in feats(m):
            a = f["areas"]
            fw.append(f"| {f['id']} | {f['name']} | {len(f['screens'])} | {len(f['apis'])} | {a['UI']} | {a['API']} | {a['Database']} | "
                      f"{a['Integration']} | {a['Testing']} | {a['Documentation']} | {f['hours']} | {f['notes']} |")
        fw += ["", f"**Module total: {m['hours']} hours**", ""]
    (ROOT / f"{CODE}_Estimation_Summary_FeatureWise.md").write_text("\n".join(fw))


# ---------------------------------------------------------------- DOCX helpers
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

NAVY = RGBColor(0x1B, 0x3A, 0x5F)


def base_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = NAVY
        doc.styles[name].font.name = "Calibri"
    return doc


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_runs(par, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = par.add_run(part.replace("`", ""))
        run.bold = i % 2 == 1


def add_table(doc, header, body, num_cols=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        shade(c, "1B3A5F")
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_runs(p, str(v))
            for r in p.runs:
                r.font.size = Pt(9.5)
            if i in num_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if ri % 2 == 1:
                shade(cells[i], "F2F4F7")
    doc.add_paragraph()
    return t


def md_to_docx(md_path: Path, out_path: Path):
    doc = base_doc()
    lines = md_path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            body = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                body.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, header, body)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2 if re.match(r"^\d{2}\. ", line[4:]) or line[4:].startswith("Appendix") else 3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(22)
            r.font.color.rgb = NAVY
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip() in ("---", ""):
            pass
        else:
            add_runs(doc.add_paragraph(), line.rstrip(" "))
        i += 1
    doc.save(out_path)


# ---------------------------------------------------------------- Step 7
def write_client_docx():
    all_h, go, opt = totals()
    doc = base_doc()
    for text, size, bold in [("Module & Feature-wise Effort Estimate", 22, True), (PROJECT, 15, True),
                             (f"Prepared for {CLIENT} · Effort in person-hours", 11, False),
                             (f"Date: {TODAY} · Project code: {CODE}", 10, False)]:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(f"This document presents a module and feature-wise effort estimate for the {PROJECT}. The platform comprises a public, "
                      "responsive online moulding catalogue integrated with the Garrison website, a secure administration portal, and a structured "
                      "catalogue data framework that lets Garrison add categories, attributes and products without software changes.")
    doc.add_paragraph("Hours cover the design and build of application software: configuration and data framework, administration and publishing "
                      "workflow, public catalogue pages, search and filters, enquiry handling, SEO and website integration, import tooling for the initial "
                      "catalogue, security, environments and deployment, with associated testing and documentation.")
    doc.add_paragraph("Hours do not cover project management, the UX/UI design phase, initial catalogue population labour (priced per product), hosting "
                      "and third-party service charges, independent penetration testing, or annual support and maintenance.")

    doc.add_heading("2. Summary Totals", level=1)
    add_table(doc, ["Bucket", "Person-hours"], [["Recommended go-live", go], ["Optional", opt], ["**Total (all items)**", f"**{all_h}**"],
                                                ["Modules", len(MODULES)], ["Features", len(FEATURES)]], num_cols=(1,))

    doc.add_heading("3. Module-wise Summary", level=1)
    add_table(doc, ["Module", "Type", "Features", "Hours", "% of total"],
              [[f"{m['nn']}. {m['name']}", m["type"], len(feats(m)), m["hours"], pct(m["hours"], all_h)] for m in MODULES]
              + [["**Total**", "", f"**{len(FEATURES)}**", f"**{all_h}**", "**100%**"]], num_cols=(2, 3, 4))

    doc.add_heading("4. Feature-wise Detail", level=1)
    for m in MODULES:
        doc.add_heading(f"{m['nn']}. {m['name']}", level=2)
        doc.add_paragraph(m["summary"] + ".")
        add_table(doc, ["Feature ID", "Feature Name", "Hours", "Notes"],
                  [[f["id"], f["name"], f["hours"], "Recommended for go-live" if f["notes"] == "Go-live" else "Optional (recommended)"] for f in feats(m)]
                  + [["", "**Module total**", f"**{m['hours']}**", ""]], num_cols=(2,))

    doc.add_heading("5. Optional Items and Commercial Line Items", level=1)
    add_table(doc, ["Item", "Basis"], [
        ["F12-004 Catalogue Export & Bulk Update", f"{opt} person-hours (optional software feature)"],
        ["Initial catalogue population — up to the agreed product count", "One-time fixed price, based on the product count confirmed at discovery"],
        ["Additional product population beyond the agreed quantity", "Per product"],
        ["Support and maintenance (AMC)", "Annual"],
    ])

    doc.add_heading("6. Effort Profile", level=1)
    areas = area_sum(FEATURES)
    label = {"UI": "User interface (screens and pages)", "API": "Application services and business logic", "Database": "Data model and storage",
             "Integration": "Integrations and environments", "Testing": "Testing", "Documentation": "Documentation"}
    add_table(doc, ["Work area", "Hours", "% of total"], [[label[a], h, pct(h, all_h)] for a, h in areas.items()], num_cols=(1, 2))

    doc.add_heading("7. Indicative Delivery Phasing", level=1)
    for t in ["Foundation: administrator access, taxonomy and masters, attribute and dimension framework, environments and release automation.",
              "Catalogue core: product profiles, media and documents, combinations and relationships, completeness rules and the review and publishing workflow.",
              "Population tooling: import templates, bulk media matching, staging and commit, running in parallel with catalogue data preparation.",
              "Customer experience: catalogue home, listings, product and combination pages, search and filters, enquiry.",
              "Go-live readiness: SEO, website integration, analytics, security hardening, backup and monitoring, testing and acceptance."]:
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("8. Items Not Included in Feature Hours", level=1)
    for t in ["Project management, discovery workshops and the UX/UI design phase",
              "Initial catalogue population labour (data extraction, normalisation, drawing preparation) — priced per product",
              "Hosting, CDN, email, analytics and search service subscriptions",
              "Obtaining third-party content rights or permissions",
              "Independent penetration testing",
              "E-commerce, pricing, inventory, customer accounts, CRM or ERP integration",
              "Multilingual content",
              "Annual support and maintenance after the warranty period"]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("9. Key Planning Assumptions", level=1)
    for t in ["The platform is a digital product catalogue; visitors do not log in.",
              "Administrator access uses configurable roles: Super Administrator, Catalogue Administrator, Catalogue Reviewer and Content Administrator.",
              "Garrison provides website and hosting access, brand assets, catalogue source files and service accounts (email, analytics, spam protection) on time.",
              "The catalogue integrates with the Garrison website through a sub-path or sub-domain confirmed during discovery.",
              "A dedicated search capability is available within the agreed hosting environment.",
              "Garrison approves taxonomy, external product inclusion and the initial dataset, and confirms third-party content rights.",
              "Content is in English; dimensions are primarily in inches with optional millimetre display.",
              "The initial product count is agreed during discovery."]:
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("10. Disclaimer", level=1)
    doc.add_paragraph("Estimates are indicative for planning and commercial discussion. Final effort may vary with the confirmed website platform and "
                      "hosting environment, the agreed initial product count, the quality of source catalogue material, and any scope changes agreed "
                      "during discovery.")
    doc.save(ROOT / f"{CODE}_Estimation_Summary_FeatureWise_Client.docx")


def main():
    validate()
    write_features()
    write_modules_md()
    write_registry()
    write_summaries()
    md_to_docx(ROOT / f"SoW_{CODE}_Detailed.md", ROOT / f"SoW_{CODE}_Detailed.docx")
    write_client_docx()
    all_h, go, opt = totals()
    print(f"Features {len(FEATURES)} · go-live {go} · optional {opt} · all-in {all_h}")
    for m in MODULES:
        print(f"  {m['nn']} {m['name']}: {len(feats(m))} features, {m['hours']} h")


if __name__ == "__main__":
    main()
