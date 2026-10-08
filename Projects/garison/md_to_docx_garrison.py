#!/usr/bin/env python3
"""Convert Garrison_Technical_Proposal.md to a formatted Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "Garrison_Technical_Proposal.md"
DOCX_PATH = ROOT / "Garrison_Technical_Proposal.docx"

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
OAK = RGBColor(0xB0, 0x7A, 0x45)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1B3A5F"
ALT_ROW = "F3F6FA"
FIGURE_BG = "FBF4EC"
NOTE_BG = "FFF4E5"

COVER_KEYS = ("Prepared for", "Prepared by", "Document type", "Version", "Date", "Proposal reference", "Classification")


def set_run_font(run, *, size=11, bold=False, italic=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_runs(paragraph, text: str, *, size=11, color=None, italic=False):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = text.replace("`", "")
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_run_font(paragraph.add_run(part[2:-2]), size=size, bold=True, italic=italic, color=color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_run_font(paragraph.add_run(part[1:-1]), size=size, italic=True, color=color)
        else:
            set_run_font(paragraph.add_run(part), size=size, italic=italic, color=color)


def shade(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def borders(table, color="CBD5E1", size="4", style="single"):
    tbl_pr = table._tbl.tblPr
    el_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        el_borders.append(el)
    tbl_pr.append(el_borders)


def spacing(p, before=0, after=6, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def hr(doc):
    p = doc.add_paragraph()
    spacing(p, 4, 10)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1B3A5F")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = lines[i].strip()
        if re.match(r"^\|[\s\-|:]+\|$", row):
            i += 1
            continue
        rows.append([c.strip() for c in row.strip("|").split("|")])
        i += 1
    return rows, i


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table)
    header_blank = all(not c for c in rows[0])
    for r, row in enumerate(rows):
        for c in range(cols):
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            spacing(p, 2, 2, 1.05)
            val = row[c] if c < len(row) else ""
            is_header = r == 0 and not header_blank
            add_runs(p, val, size=9.5 if is_header else 10, color=WHITE if is_header else None)
            if val in ("✓", "—"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if is_header:
                shade(cell, HEADER_BG)
                for run in p.runs:
                    run.bold = True
            elif r % 2 == 0:
                shade(cell, ALT_ROW)
    doc.add_paragraph()


def add_callout(doc, block_lines):
    is_note = block_lines and "EDITOR'S NOTE" in block_lines[0]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(table, color="B07A45", size="8", style="dashed")
    cell = table.rows[0].cells[0]
    shade(cell, NOTE_BG if is_note else FIGURE_BG)
    cell.text = ""
    first = True
    for line in block_lines:
        if not line.strip():
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        spacing(p, 2, 3, 1.1)
        if first:
            add_runs(p, line, size=11, color=OAK)
            for run in p.runs:
                run.bold = True
            first = False
        else:
            add_runs(p, line, size=9.5)
    doc.add_paragraph()


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    sec.left_margin = sec.right_margin = Inches(0.95)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11.5, OAK)):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def footer(doc, text):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_run_font(run, size=8, color=MUTED)


def cover(doc, meta):
    for _ in range(4):
        doc.add_paragraph()
    for text, size, color, bold in (
        ("TECHNICAL PROPOSAL", 13, OAK, True),
        ("Online Moulding Catalogue Platform", 24, NAVY, True),
        (f"Prepared for {meta.get('Prepared for', 'Garrison Industries')}", 13, MUTED, False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, 6, 6)
        set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    hr(doc)
    for key in COVER_KEYS[1:]:
        if key in meta:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(p, 1, 1)
            set_run_font(p.add_run(f"{key}: "), size=10.5, bold=True, color=MUTED)
            set_run_font(p.add_run(meta[key]), size=10.5, color=MUTED)
    doc.add_page_break()


def convert(md: str) -> Document:
    lines = md.splitlines()
    doc = Document()
    configure(doc)

    meta = {}
    body_start = 0
    for idx, line in enumerate(lines):
        m = re.match(r"\*\*(.+?):\*\*\s*(.+)", line.strip())
        if m and m.group(1) in COVER_KEYS:
            meta[m.group(1)] = m.group(2).strip()
        if line.strip() == "---" and meta:
            body_start = idx + 1
            break
    cover(doc, meta)

    i = body_start
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s == "---":
            i += 1
            continue
        if s.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            add_callout(doc, block)
            continue
        if s.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if s.startswith("### "):
            spacing(doc.add_heading(s[4:], level=3), 10, 4)
            i += 1
            continue
        if s.startswith("## "):
            title = s[3:]
            if title.startswith(("2. ", "16. ", "31. ")):
                doc.add_page_break()
            spacing(doc.add_heading(title, level=1), 14, 8)
            i += 1
            continue
        if s.startswith("# "):
            i += 1
            continue
        if re.match(r"^\*Figure \d+", s):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(p, 0, 10)
            set_run_font(p.add_run(s.strip("*")), size=9.5, italic=True, color=MUTED)
            i += 1
            continue
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            spacing(p, 0, 3)
            add_runs(p, s[2:])
            i += 1
            continue
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            spacing(p, 0, 3)
            add_runs(p, re.sub(r"^\d+\.\s+", "", s))
            i += 1
            continue
        p = doc.add_paragraph()
        spacing(p, 2, 6)
        if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, s, color=MUTED)
        else:
            add_runs(p, s)
        i += 1

    footer(doc, "InvenTech Info Solutions  |  Garrison Industries — Online Moulding Catalogue Proposal  |  Confidential  |  Page ")
    return doc


def main():
    convert(MD_PATH.read_text(encoding="utf-8")).save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()
