"""Build draft security policy documents (Word) for HDFC vendor onboarding.

Run from repo root:
    .venv/bin/python Projects/hdfc_enrollment/_build/build_policies.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_checklist_docs import NAVY, VENDOR, table  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "Policies"
DATE = "9 October 2026"
TBF = "[To be filled]"

# ---------------------------------------------------------------------------
# Policy content. Blocks: ("h", text) | ("p", text) | ("b", [bullets]) |
# ("t", headers, rows, widths)
# ---------------------------------------------------------------------------

RACI = ("t", ["Sl. No.", "Activity", "Mgmt", "ISO", "PM", "Dev", "DevOps", "Staff"], [
    (1, "Approve this policy and provide resources", "A", "R", "I", "I", "I", "I"),
    (2, "Classify project data", "I", "C", "A/R", "C", "I", "–"),
    (3, "Grant, review and remove user access", "I", "A", "R", "I", "R", "–"),
    (4, "Build security and encryption into applications", "–", "C", "A", "R", "C", "–"),
    (5, "Manage servers, certificates, keys and backups", "–", "A", "I", "C", "R", "–"),
    (6, "Report security incidents", "I", "A", "R", "R", "R", "R"),
    (7, "Conduct awareness training and keep records", "I", "A/R", "C", "R", "R", "R"),
    (8, "Review this policy every year", "A", "R", "C", "C", "C", "I"),
], [1.2, 6.2, 1.6, 1.6, 1.6, 1.6, 1.6, 1.6])

RACI_LEGEND = ("p", "Mgmt: Management. ISO: Information Security Officer. PM: Project Manager. Dev: Development team. "
                    "DevOps: DevOps / Infrastructure team. Staff: all employees and contractors. "
                    "R: Responsible. A: Accountable. C: Consulted. I: Informed.")

TRAINING_RECORDS = ("b", [
    "Every employee and contractor completes security and data protection awareness training within 30 days of joining, "
    "and a refresher once a year.",
    "Training records are kept for at least three years. Each record shows the date, topic, trainer, list of attendees "
    "(signed attendance sheet or screenshot from the online session) and assessment result where a test is given.",
    "The Information Security Officer reviews completion every quarter and follows up with anyone who has not completed "
    "the training.",
])


def review(n):
    return [("h", f"{n}. Compliance, exceptions and review"),
            ("b", ["Breach of this policy may lead to disciplinary action, in line with the Company's HR policy.",
                   "Any exception must be approved in writing by the Information Security Officer, with a reason and an "
                   "expiry date, and recorded in the exceptions register.",
                   "This policy is reviewed at least once a year, and also after a major security incident, a change in "
                   "law or a major change in the Company's systems."])]


POLICIES = [
    dict(
        file="IIS_Data_Protection_Policy.docx",
        title="Data Protection Policy",
        doc_id="IIS-POL-DP-01",
        blocks=[
            ("h", "1. Purpose"),
            ("p", "This policy sets out how Inventech Info Solutions Pvt. Ltd. (\"the Company\") protects the data it "
                  "handles, both its own data and data it processes for clients, against loss, misuse, and unauthorised "
                  "access or change."),
            ("h", "2. Scope"),
            ("b", ["All employees, trainees and contractors of the Company.",
                   "All systems used to build, test, host or support applications: laptops, source code repositories, the "
                   "Company's own UAT servers, and production environments run for clients (for example at the Karnataka "
                   "State Data Centre or on client premises). Where a client hosts the system, the Company follows this "
                   "policy for the parts it controls and the client's policy where that is stricter.",
                   "All data in any form: databases, files, documents, backups, logs, email and paper."]),
            ("h", "3. References"),
            ("b", ["Digital Personal Data Protection Act, 2023.",
                   "Information Technology Act, 2000, and the IT (Reasonable Security Practices and Procedures and "
                   "Sensitive Personal Data or Information) Rules, 2011.",
                   "CERT-In Directions under Section 70B(6) of the IT Act, dated 28 April 2022.",
                   "ISO/IEC 27001:2022 (used as guidance).",
                   "Security requirements in client contracts."]),
            ("h", "4. Data classification schema"),
            ("p", "All data is placed in one of four classes. The Project Manager classifies project data at the start of "
                  "each project. Data that has not been classified is treated as Confidential."),
            ("t", ["Sl. No.", "Class", "Description", "Examples", "Handling"], [
                (1, "Public", "Approved for public release.", "Published website content, public notices.",
                 "No limit on disclosure. Protect against unauthorised change."),
                (2, "Internal", "Meant for Company staff only.", "Internal procedures, project plans, general training "
                 "material.", "Share only with staff and authorised contractors. Do not post publicly."),
                (3, "Confidential", "Business or client data whose disclosure could cause harm.", "Client requirements "
                 "and designs, source code, architecture, configuration, commercial proposals.", "Need-to-know access. "
                 "Encrypt in transit. Store only on Company-approved systems."),
                (4, "Restricted", "Personal data and secrets. Disclosure could cause serious harm to people or "
                 "clients.", "Names, mobile numbers, addresses, Aadhaar numbers and identity documents of citizens or "
                 "candidates; payroll data; passwords, API keys, encryption keys and private keys.",
                 "Access only with data owner approval. Encrypt in transit and at rest. Never place in source code, "
                 "email, chat or test environments. Mask in logs and reports."),
            ], [1.2, 2.2, 3.6, 5, 5]),
            ("h", "5. Roles and responsibilities (RACI matrix)"),
            RACI, RACI_LEGEND,
            ("h", "6. Access control"),
            ("b", ["Access is given on a need-to-know basis with the least privilege needed for the job. Applications use "
                   "role-based access control (roles, permissions and role–permission mapping).",
                   "Every user has a unique user ID. Shared accounts are not allowed.",
                   "Multi-factor authentication (OTP) is used for administrator access to applications and for access to "
                   "servers, source code repositories and cloud consoles wherever the system supports it.",
                   "Passwords have at least 10 characters with upper case, lower case, a number and a special character. "
                   "Accounts lock after 5 failed attempts. Passwords are stored only as salted hashes (bcrypt).",
                   "The Project Manager approves access requests. Production access also needs approval from the "
                   "Information Security Officer.",
                   "Access is removed within one working day when a person leaves or changes role.",
                   "All access rights are reviewed every quarter and the review is recorded.",
                   "Vendors and third parties are not given access to Company systems or client applications. Any future "
                   "need requires a signed contract and non-disclosure agreement, a named account with an expiry date, and "
                   "written approval from the Project Manager and the Information Security Officer."]),
            ("h", "7. Encryption requirements"),
            ("t", ["Sl. No.", "Data", "Requirement"], [
                (1, "Data in transit (users, APIs, integrations)", "TLS 1.3 with certificates from a trusted certificate "
                 "authority. SSL and TLS 1.0 / 1.1 are disabled. Details are in the Data Transmission Policy."),
                (2, "Data at rest: databases, file and object storage, and backups that hold Confidential or Restricted "
                 "data", "AES-256 encryption."),
                (3, "Passwords", "Salted one-way hash (bcrypt). Never stored or logged in plain text."),
                (4, "OTPs and session tokens", "Stored only as SHA-256 hashes, with a short expiry."),
                (5, "Login tokens", "Signed with HMAC-SHA256 or stronger, with a short expiry."),
                (6, "Digital certificates", "Validity of one year or less, SHA-256 signature, renewed before expiry."),
                (7, "Keys and secrets", "Kept outside source code in environment-specific secret stores. Access limited "
                 "to the DevOps team. Rotated at least once a year, and at once if exposed."),
            ], [1.2, 6, 9.8]),
            ("h", "8. Data in development and testing"),
            ("b", ["Production personal data is not copied to development or UAT environments. Test data is synthetic or "
                   "masked.",
                   "Source code repositories must not contain personal data, passwords, keys or production configuration. "
                   "Repositories are scanned for secrets before code is merged.",
                   "Every code change is reviewed through a pull request before it is merged.",
                   "Logs and audit records mask passwords, OTPs, tokens and personal identifiers."]),
            ("h", "9. Backup, retention and disposal"),
            ("b", ["Production backups follow the schedule agreed with the client or hosting provider.",
                   "For the Company's own servers, databases are backed up daily, backups are kept for 30 days, and a "
                   "restore test is done every quarter.",
                   "Data is kept only as long as the client contract or the law requires. At the end of a contract, data is "
                   "returned or deleted as the client instructs, and the deletion is recorded.",
                   "Storage media are wiped securely before reuse. Media that cannot be wiped are physically destroyed, "
                   "and a disposal certificate is kept."]),
            ("h", "10. Awareness and training"),
            TRAINING_RECORDS,
            *review(11),
        ]),
    dict(
        file="IIS_Incident_Response_Plan.docx",
        title="Incident Response Plan",
        doc_id="IIS-POL-IR-01",
        blocks=[
            ("h", "1. Purpose"),
            ("p", "This plan explains how the Company detects, responds to and recovers from information security "
                  "incidents. The aim is to limit damage, restore services quickly and meet our reporting duties to "
                  "clients and regulators."),
            ("h", "2. Scope"),
            ("p", "The plan covers incidents that affect the Company's own systems, the applications the Company builds or "
                  "supports for clients (including those hosted at the Karnataka State Data Centre or on client premises), "
                  "and any data the Company processes."),
            ("h", "3. Definitions"),
            ("t", ["Sl. No.", "Term", "Meaning"], [
                (1, "Security event", "Something observed in a system or network that may point to a security problem, "
                 "for example repeated failed logins."),
                (2, "Security incident", "An event that harms, or is likely to harm, the confidentiality, integrity or "
                 "availability of data or systems."),
                (3, "Personal data breach", "Any unauthorised processing, loss, disclosure or change of personal data, "
                 "as defined in the Digital Personal Data Protection Act, 2023."),
            ], [1.2, 4, 11.8]),
            ("h", "4. Incident response team"),
            ("t", ["Sl. No.", "Role", "Responsibilities", "Name and contact"], [
                (1, "Incident Manager (Information Security Officer)", "Owns each incident from start to finish. Decides "
                 "severity, leads the team, and reports to management, clients and CERT-In.", TBF),
                (2, "Technical Lead", "Investigates the cause in the application. Prepares and deploys the fix.", TBF),
                (3, "DevOps / Infrastructure Lead", "Contains the incident at server and network level: isolates "
                 "servers, blocks IP addresses, changes credentials, restores from backup. Preserves logs.", TBF),
                (4, "Project Manager", "Single point of contact with the client. Shares updates on the agreed timeline.",
                 TBF),
                (5, "Management representative", "Approves external communication and decisions with business or legal "
                 "impact.", TBF),
                (6, "Legal / compliance advisor", "Advises on duties under the DPDP Act, the IT Act and client "
                 "contracts.", TBF),
            ], [1.2, 4.2, 8, 3.6]),
            ("h", "5. Incident classification and severity"),
            ("t", ["Sl. No.", "Severity", "Description", "Examples", "Acknowledge within", "Contain within"], [
                (1, "S1 Critical", "Confirmed breach of Restricted data, or a production system down for all users.",
                 "Leak of citizen personal data; ransomware; defaced public portal; compromised administrator account.",
                 "15 minutes", "4 hours"),
                (2, "S2 High", "Serious threat to Confidential or Restricted data, or a major function unavailable.",
                 "Malware on a server; suspected unauthorised access; exposed credentials or keys.", "1 hour",
                 "24 hours"),
                (3, "S3 Medium", "Limited impact and no confirmed data exposure.", "Password-guessing attack on a login "
                 "page; vulnerability reported in production; lost laptop with an encrypted disk.", "4 hours",
                 "3 working days"),
                (4, "S4 Low", "Minor policy breach or suspicious event.", "Phishing email reported and not clicked; "
                 "password shared within the team.", "1 working day", "7 working days"),
            ], [1.2, 2, 4.2, 5.2, 2.2, 2.2]),
            ("h", "6. Escalation matrix"),
            ("t", ["Sl. No.", "Severity", "First contact", "Escalate to", "Then to", "Client informed"], [
                (1, "S1 Critical", "Incident Manager, at once", "Management representative, within 1 hour",
                 "Managing Director", "Within 2 hours of confirmation, or as per contract"),
                (2, "S2 High", "Incident Manager, within 1 hour", "Project Manager and Technical Lead",
                 "Management representative, if not contained in 24 hours",
                 "Within 24 hours if client data or service is affected"),
                (3, "S3 Medium", "Technical Lead", "Incident Manager", "–", "In the monthly report, or as agreed"),
                (4, "S4 Low", "Reporting person's manager", "Incident Manager", "–", "Not required"),
            ], [1.2, 2, 3.2, 3.6, 3.4, 3.6]),
            ("h", "7. Detection and reporting"),
            ("p", "7.1 Sources of detection"),
            ("b", ["Application logs and audit trails (logins, data changes, administrator actions).",
                   "Login history and alerts on repeated failed logins.",
                   "Server and infrastructure monitoring.",
                   "Alerts from hosting providers (Karnataka State Data Centre, client IT teams).",
                   "Reports from clients, users, employees and vendors."]),
            ("p", "7.2 Internal reporting"),
            ("b", ["Anyone who notices or suspects an incident reports it at once to the Incident Manager by phone and by "
                   f"email ({TBF}).",
                   "Do not investigate on your own, and do not delete or change anything that may be evidence."]),
            ("p", "7.3 External reporting"),
            ("b", ["CERT-In: within 6 hours of noticing the incident, for the incident types listed in the CERT-In "
                   "Directions of 28 April 2022.",
                   "Client: as per the contract. Where the contract is silent, as per the escalation matrix above.",
                   "Personal data breach: the client, as data fiduciary, is informed without delay so that it can notify the "
                   "Data Protection Board of India and affected persons under the DPDP Act, 2023. The Company gives the "
                   "client all the information it needs."]),
            ("h", "8. Standard operating procedure (SOP)"),
            ("t", ["Sl. No.", "Step", "Actions", "Owner"], [
                (1, "Identify and log", "Record the incident in the incident register, assign severity and inform the "
                 "team.", "Incident Manager"),
                (2, "Contain", "Isolate affected systems, disable compromised accounts, block malicious IP addresses and "
                 "change exposed credentials or keys. Preserve evidence (logs, snapshots) before making changes.",
                 "DevOps / Infrastructure Lead"),
                (3, "Investigate", "Find the cause and the extent: which data, which users and which time period.",
                 "Technical Lead"),
                (4, "Eradicate", "Remove malware, fix the vulnerability and apply patches.",
                 "Technical Lead, DevOps Lead"),
                (5, "Recover", "Restore from a clean backup, verify the system and watch it closely for 7 days.",
                 "DevOps / Infrastructure Lead"),
                (6, "Communicate", "Update the client and management as per the escalation matrix, and make external "
                 "reports.", "Project Manager, Incident Manager"),
                (7, "Close and review", "For S1 and S2 incidents, complete a root cause analysis (RCA) report within 5 "
                 "working days of closure. Record lessons learnt and update controls and this plan.",
                 "Incident Manager"),
            ], [1.2, 3, 9, 3.8]),
            ("h", "9. Incident register and evidence"),
            ("b", ["Every incident is recorded in the incident register with: incident ID, date and time detected, "
                   "reported by, system affected, severity, description, actions taken, data affected, notifications "
                   "made (to whom and when), root cause and closure date.",
                   f"The incident register is maintained in {TBF} (tool or format).",
                   "Logs of all systems are kept for 180 days within India, as required by the CERT-In Directions.",
                   "Evidence and RCA reports are kept for at least one year after the incident is closed."]),
            ("h", "10. Training and awareness"),
            ("b", ["All staff learn how to recognise and report incidents during induction and in the yearly refresher.",
                   "The incident response team runs a tabletop exercise at least once a year.",
                   "Training and exercise records are kept for at least three years."]),
            *review(11),
        ]),
    dict(
        file="IIS_Data_Transmission_Policy.docx",
        title="Data Transmission Policy",
        doc_id="IIS-POL-DT-01",
        blocks=[
            ("h", "1. Purpose"),
            ("p", "This policy makes sure that data moving between users, systems and third parties is protected against "
                  "interception and tampering."),
            ("h", "2. Scope"),
            ("p", "It covers every transfer of Company or client data, including:"),
            ("b", ["Browsers and mobile apps to application servers.",
                   "Application servers to databases, file or object storage and caches.",
                   "Integrations with third-party services such as SMS gateways, payment gateways, email and digital "
                   "signature services.",
                   "Transfers between environments (deployments and backups).",
                   "File exchange with clients and vendors, email and messaging."]),
            ("h", "3. Encryption standards"),
            ("t", ["Sl. No.", "Channel", "Minimum standard"], [
                (1, "Web applications, mobile apps and APIs", "HTTPS with TLS 1.3. TLS 1.2 is allowed only with a "
                 "written exception, where client or third-party infrastructure does not support TLS 1.3. SSL and TLS "
                 "1.0 / 1.1 are disabled. HTTP Strict Transport Security (HSTS) is enabled."),
                (2, "SSL / TLS certificates", "Issued by a trusted certificate authority. RSA 2048-bit or ECDSA P-256 "
                 "keys or stronger. SHA-256 signature. Validity of one year or less. Recorded in a certificate register "
                 "and renewed at least 15 days before expiry."),
                (3, "Application to database, storage and cache", "Encrypted (TLS) whenever traffic leaves the server. "
                 "Otherwise limited to a private network."),
                (4, "Third-party integrations (SMS, payment, email, digital signature)", "HTTPS only. Credentials and "
                 "API keys are kept on the server in secret stores. Requests are signed or hashed where the provider "
                 "supports it (for example HMAC-SHA256 or SHA-512), and responses are verified (for example the payment "
                 "gateway hash check)."),
                (5, "File transfer and backups", "SFTP or HTTPS. Backup files are encrypted with AES-256 before they "
                 "leave the server."),
                (6, "Email", "TLS between mail servers. Restricted data is sent only as a password-protected (AES-256) "
                 "attachment, with the password shared on a separate channel."),
                (7, "Removable media", "Not used for Restricted data. If unavoidable, the data is AES-256 encrypted and "
                 "the Information Security Officer approves it."),
                (8, "Data at rest (for reference)", "AES-256, as set out in the Data Protection Policy."),
            ], [1.2, 4.6, 11.2]),
            ("h", "4. Rules for development and operations"),
            ("b", ["No personal data, passwords or tokens in URLs or query strings.",
                   "Session cookies are marked Secure, HttpOnly and SameSite.",
                   "SMS messages carry only the information needed (for example an OTP or a status update).",
                   "Cross-origin access (CORS) is limited to known domains.",
                   "Certificate checks are never switched off in code, and production mobile builds do not allow "
                   "unencrypted traffic.",
                   "The TLS setup is tested before go-live and after every certificate renewal, and the test result is kept "
                   "as evidence."]),
            ("h", "5. Roles and responsibilities"),
            ("t", ["Sl. No.", "Role", "Responsibility"], [
                (1, "Information Security Officer", "Owns this policy, approves exceptions and checks compliance."),
                (2, "DevOps / Infrastructure team", "Configures TLS and certificates, keeps the certificate register and "
                 "renews certificates on time."),
                (3, "Development team", "Builds integrations as per this policy and keeps secrets out of code."),
                (4, "Project Manager", "Makes sure client-hosted environments meet this policy, or records an "
                 "exception."),
            ], [1.2, 5, 10.8]),
            ("h", "6. Training and awareness"),
            ("b", ["Data transmission rules are part of the yearly security awareness training for all staff.",
                   "Developers and the DevOps team attend a secure integration session once a year.",
                   "Training records are kept for at least three years."]),
            *review(7),
        ]),
    dict(
        file="IIS_Data_Privacy_Policy.docx",
        title="Data Privacy Policy",
        doc_id="IIS-POL-PR-01",
        blocks=[
            ("h", "1. Purpose"),
            ("p", "This policy explains how the Company collects, uses, shares and protects personal data, in line with "
                  "the Digital Personal Data Protection Act, 2023 (DPDP Act)."),
            ("h", "2. Scope"),
            ("b", ["Personal data processed on behalf of clients in the applications the Company builds, hosts or "
                   "supports. Here the client is the data fiduciary and the Company is the data processor.",
                   "Personal data of the Company's own employees, job applicants and business contacts. Here the Company "
                   "is the data fiduciary."]),
            ("h", "3. Principles"),
            ("b", ["Lawful and fair: personal data is processed only for a lawful purpose.",
                   "Purpose limitation: data is used only for the purpose it was collected for.",
                   "Data minimisation: only the data needed for that purpose is collected.",
                   "Accuracy: data is kept accurate and can be corrected.",
                   "Storage limitation: data is kept only as long as needed.",
                   "Security: data is protected as set out in the Data Protection Policy.",
                   "Accountability: the Company can show how it meets this policy."]),
            ("h", "4. Data collection: what data is collected and why"),
            ("t", ["Sl. No.", "Category", "Examples", "Why it is collected"], [
                (1, "Contact details of citizens and applicants", "Name, mobile number, email, address",
                 "Registration, OTP login and status updates by SMS, as the client's service requires."),
                (2, "Identity and supporting documents", "Identity proof, BPL certificate, institution registration "
                 "documents", "To check eligibility as defined in the client's process."),
                (3, "Candidate and employee profile data", "Employee ID, department, photograph, Aadhaar number (only "
                 "where the client requires it)", "Exam registration, identity check at the exam centre and "
                 "certificates."),
                (4, "Field and location data", "Site photos, video, GPS location", "Verification of site dimensions."),
                (5, "Payment information", "Transaction reference, amount, status", "Confirmation of fee payment. Card "
                 "and bank details are entered only on the payment gateway's page and are not stored in our "
                 "applications."),
                (6, "Client staff user data", "Name, official email, mobile number, role, login history",
                 "Access control and audit trail."),
                (7, "Payroll and time-study data of client employees", "Employee names, time records, salary cost",
                 "Administrative cost claiming, as per the client contract."),
                (8, "Company employees and job applicants", "Identity and address proof, bank details, qualifications",
                 "HR, payroll and statutory compliance."),
            ], [1.2, 4, 5.4, 6.4]),
            ("b", ["Only the fields in the client's approved requirements are collected. Any new personal data field needs "
                   "approval from the Project Manager and the client.",
                   "Aadhaar numbers are collected only where the law or the client requires it. They are masked on screen "
                   "(only the last four digits are shown) and encrypted at rest."]),
            ("h", "5. Transparency and communication"),
            ("b", ["Every public-facing application shows a privacy notice, linked from the registration and login pages. "
                   "It is written in plain language, and in English and Kannada where the application is bilingual.",
                   "The notice says what data is collected, why, how long it is kept, who it is shared with, how to ask "
                   "for correction or deletion, and whom to contact with a grievance.",
                   "Where consent is the legal basis, the application records consent with a date and time, together with "
                   "a link to the notice that was shown.",
                   "For client applications, the client (as data fiduciary) approves the content of the notice."]),
            ("h", "6. Rights of individuals"),
            ("b", ["Individuals can ask to access, correct or erase their personal data, nominate another person, and raise "
                   "a grievance.",
                   "For client applications, requests received by the Company are passed to the client within two working "
                   "days, and the Company gives the technical help needed to act on them.",
                   "For the Company's own data, the Grievance Officer responds within 30 days."]),
            ("h", "7. Third parties and vendors"),
            ("p", "Personal data is shared only with third parties needed to deliver the service, and only under a written "
                  "contract. Each contract covers:"),
            ("b", ["Confidentiality, and use of the data only for the agreed purpose.",
                   "Security safeguards at least equal to the Company's policies.",
                   "Notice to the Company of any breach within 24 hours.",
                   "No sub-processing without the Company's written approval.",
                   "Where the data is stored.",
                   "Return or deletion of the data when the contract ends.",
                   "The Company's right to audit."]),
            ("p", "Third-party services currently used in client applications:"),
            ("t", ["Sl. No.", "Service", "Purpose", "Data shared"], [
                (1, "Karnataka Mobile Seva SMS gateway (Government of Karnataka)", "OTP and status SMS",
                 "Mobile number, message text"),
                (2, "Payment gateways (PayU, Razorpay)", "Online fee payment",
                 "Name, mobile number or email, amount, transaction reference"),
                (3, "Hosting: Karnataka State Data Centre and client data centres", "Production hosting, arranged by the "
                 "client", "All application data"),
            ], [1.2, 6, 4.4, 5.4]),
            ("p", "This list is kept up to date in the Company's third-party register."),
            ("h", "8. Retention and deletion"),
            ("b", ["Personal data is kept only as long as the purpose, the client contract or the law requires.",
                   "When a contract ends, data is returned or deleted as the client instructs, and the deletion is "
                   "recorded."]),
            ("h", "9. Data location"),
            ("b", ["Personal data in applications for Indian clients is hosted in India (at the Karnataka State Data Centre "
                   "or in client data centres).",
                   "For overseas clients, data is hosted on the client's own infrastructure, as the client instructs.",
                   "Personal data is transferred outside India only where the DPDP Act and the client contract allow it."]),
            ("h", "10. Security and breaches"),
            ("p", "Personal data is protected as set out in the Data Protection Policy and the Data Transmission Policy. "
                  "Personal data breaches are handled under the Incident Response Plan."),
            ("h", "11. Roles and responsibilities"),
            ("t", ["Sl. No.", "Role", "Responsibility", "Name and contact"], [
                (1, "Grievance Officer / Data Protection contact", "Handles privacy requests and grievances. Advises on the "
                 "DPDP Act.", TBF),
                (2, "Information Security Officer", "Owns this policy and checks compliance.", TBF),
                (3, "Project Manager", "Makes sure each application collects only approved data and shows the privacy "
                 "notice.", "–"),
                (4, "All staff", "Handle personal data as set out in this policy, and report any misuse.", "–"),
            ], [1.2, 4.6, 7.6, 3.6]),
            ("h", "12. Training and awareness"),
            TRAINING_RECORDS,
            *review(13),
        ]),
]

# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def page_field(run):
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)


def new_document(policy):
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size in (("Heading 1", 13), ("Heading 2", 11)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(NAVY)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("DRAFT – for management review")
    r.font.size = Pt(8)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("C00000")

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{VENDOR} | {policy['doc_id']} | Internal | Page ").font.size = Pt(8)
    run = footer.add_run()
    run.font.size = Pt(8)
    page_field(run)

    for text, size, color in ((VENDOR, 11, "595959"), (policy["title"], 20, NAVY)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)

    doc.add_heading("Document control", level=1)
    table(doc, ["Item", "Details"], [
        ("Document title", policy["title"]),
        ("Document ID", policy["doc_id"]),
        ("Version", "0.1 (draft for management review)"),
        ("Classification", "Internal"),
        ("Document owner", policy.get("owner", "Information Security Officer")),
        ("Approved by", TBF),
        ("Effective date", "Date of approval"),
        ("Next review", "12 months from the effective date"),
    ], [5, 12])
    table(doc, ["Sl. No.", "Version", "Date", "Description of change", "Prepared by"],
          [(1, "0.1", DATE, "First draft", TBF)], [1.2, 1.8, 3.2, 6.8, 4])
    return doc


def render(doc, blocks):
    for block in blocks:
        kind = block[0]
        if kind == "h":
            doc.add_heading(block[1], level=1)
        elif kind == "p":
            doc.add_paragraph(block[1])
        elif kind == "b":
            for item in block[1]:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "t":
            table(doc, block[1], block[2], block[3])


def approval(doc):
    doc.add_heading("Approval", level=1)
    table(doc, ["Sl. No.", "Role", "Name", "Designation", "Signature", "Date"],
          [(1, "Prepared by", "", "", "", ""), (2, "Reviewed by", "", "", "", ""), (3, "Approved by", "", "", "", "")],
          [1.2, 2.8, 3.6, 3.4, 3.4, 2.6])


def main():
    OUT.mkdir(exist_ok=True)
    for policy in POLICIES:
        doc = new_document(policy)
        render(doc, policy["blocks"])
        approval(doc)
        path = OUT / policy["file"]
        doc.save(path)
        print(f"wrote {path.relative_to(Path.cwd()) if path.is_relative_to(Path.cwd()) else path}")


if __name__ == "__main__":
    main()
