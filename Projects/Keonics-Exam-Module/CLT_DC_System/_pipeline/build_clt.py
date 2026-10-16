"""Build CLT estimation artefacts (Steps 1-3, 5, 6, 7) from module and feature data.

Run from repo root:
    .venv/bin/python Projects/Keonics-Exam-Module/CLT_DC_System/_pipeline/build_clt.py
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

from clt_data_a import FEATURES as FA  # noqa: E402
from clt_data_b import FEATURES as FB  # noqa: E402
from clt_data_c import FEATURES as FC  # noqa: E402
from clt_data_d import FEATURES as FD  # noqa: E402
from clt_data_e import FEATURES as FE  # noqa: E402
from clt_data_f import FEATURES as FF  # noqa: E402
from clt_modules import MODULES  # noqa: E402

FEATURES = FA + FB + FC + FD + FE + FF
CODE = "CLT"
PROJECT = "Online Computer Literacy Test (CLT) and Digital Certification System"
CLIENT = "Karnataka State Electronics Development Corporation Limited (KEONICS)"
END_CLIENT = "Center for E-Governance, Department of Personnel and Administrative Reforms, Government of Karnataka"
TODAY = dt.date.today().isoformat()

CX = {"S": "Simple", "E": "Easy", "M": "Medium", "C": "Complex", "V": "Very High"}
HOURS = {
    "UI": {"S": 2, "E": 4, "M": 6, "C": 12, "V": 16},
    "API": {"S": 2, "E": 4, "M": 6, "C": 8, "V": 16},
    "Database": {"S": 1, "E": 2, "M": 4, "C": 6, "V": 10},
    "Integration": {"S": 4, "E": 8, "M": 16, "C": 24, "V": 40},
    "Testing": {"S": 2, "E": 4, "M": 6, "C": 8, "V": 12},
    "Documentation": {"S": 1, "E": 2, "M": 3, "C": 4, "V": 6},
}
AREA_KEYS = [("UI", "screens"), ("API", "apis"), ("Database", "db"), ("Integration", "integ"), ("Testing", "tests"),
             ("Documentation", "docs")]
GO_LIVE = "Go-live"

def slugify(text: str) -> str:
    text = text.lower().replace("&", "and")
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


for m in MODULES:
    m["slug"] = slugify(m["name"])
    m["folder"] = f"modules/{m['nn']}-{m['slug']}"


def rows(f):
    out = []
    for area, key in AREA_KEYS:
        for name, cx in f.get(key, []):
            out.append((area, name, CX[cx], HOURS[area][cx]))
    return out


def area_totals(f):
    t = {a: 0 for a, _ in AREA_KEYS}
    for area, _, _, h in rows(f):
        t[area] += h
    return t


for f in FEATURES:
    f.setdefault("notes", GO_LIVE)
    f["client"] = "Recommended for go-live" if f["notes"] == GO_LIVE else "Conditional — taken up only if required by KEONICS"
    f["slug"] = slugify(f["name"])
    f["nn"] = f["id"][1:3]
    f["seq"] = f["id"][4:]
    f["hours"] = sum(r[3] for r in rows(f))
    f["areas"] = area_totals(f)


def module_of(f):
    return next(m for m in MODULES if m["nn"] == f["nn"])


def feats(m):
    return [f for f in FEATURES if f["nn"] == m["nn"]]


def bullets(items):
    return "\n".join(f"- {i}" for i in items) if items else "- None"


# ---------------------------------------------------------------- Step 1
def write_sow():
    all_feats = len(FEATURES)
    L = [f"# Scope of Work (Detailed) — {PROJECT}", "",
         f"**Client:** {CLIENT}  ", f"**On behalf of:** {END_CLIENT}  ",
         "**Source:** RFP KEONICS/2026-27/SE0140/Call-3 — Selection of Service Provider for Design, Development, Implementation and "
         "Maintenance of Online Computer Literacy Test (CLT) and Digital Certification System  ",
         f"**Project code:** {CODE}  ", f"**Prepared:** {TODAY}", "",
         "## 1. Introduction", "",
         "The Government of Karnataka requires every state government employee to pass the Computer Literacy Test (CLT) during probation or "
         "within the prescribed period. More than five lakh employees fall under the rule, with about 15,000 new recruits every year. KEONICS, "
         "on behalf of the Center for E-Governance (DPAR), is procuring a Service Provider to design, develop, implement and maintain an "
         "end-to-end online CLT and digital certification system hosted at the Karnataka State Data Centre (SDC).", "",
         "This document converts the RFP Terms of Reference, Annexure 1 (technical requirements) and the Special Conditions of Contract into "
         f"a structured scope of {len(MODULES)} modules and {all_feats} features for estimation and delivery planning.", "",
         "## 2. Programme Overview", "",
         "| Item | Detail |", "|------|--------|",
         "| Candidates | Government employees (with or without KGID) and prospective candidates (new appointees, board and corporation staff) |",
         "| Volume | About 25,000 registrations in 2026 and 20,000 in 2027; 15,000 new employees per year thereafter |",
         "| Test pattern | Part 1: 60 MCQs across 12 topics (60 marks). Part 2: English and Kannada typing tests (20 marks each). 90 minutes; 50% to qualify |",
         "| Sessions | Four sessions per day with an optional fifth; working Saturdays for employees and Sundays for prospective candidates |",
         "| Concurrency | At least 2,000 concurrent candidates per session |",
         "| Hosting | Karnataka State Data Centre only; no local database at test centres |",
         "| Proctoring | Continuous screen and webcam video recording with automated monitoring and real-time alerts |",
         "| Certificates | Digitally signed, watermarked, QR-verifiable certificates; results updated in HRMS |",
         "| Contract | Development and go-live by T + 3 months, 12-month warranty, then comprehensive AMC for 3 years |", "",
         "## 3. Objectives", "",
         "1. Conduct a fair, secure and fully online CLT with registration, booking, testing, evaluation and certification in one system.",
         "2. Validate identity through HRMS (KGID) or Aadhaar eKYC and at the centre through hall ticket, photo and biometric checks.",
         "3. Keep the question bank confidential through maker-checker workflow, component segregation and AES-256 encryption with HSM keys.",
         "4. Deliver unique jumbled papers per candidate with immediate scoring and tamper-proof digital certificates.",
         "5. Give KEONICS and competent authorities real-time monitoring, complete audit trails and MIS.",
         "6. Meet performance (2,000 concurrent, under one second on LAN), availability (99.9%) and Safe-to-Host requirements.", "",
         "## 4. Scope Summary", "",
         "### In scope", "",
         "- Design, development, testing, deployment at SDC and go-live of all modules listed in Section 6",
         "- Web portal at clt.karnataka.gov.in and candidate mobile application (Android and iOS)",
         "- Content Authoring Tool, question bank workflow and question paper generation",
         "- Test centre console, secure exam client, online examination engine with typing tests",
         "- Screen and webcam video proctoring with automated monitoring",
         "- Integrations: HRMS, Aadhaar authentication and eKYC, payment gateway, SMS and e-mail, HSM, DSC / e-Sign",
         "- Porting of legacy CLT data",
         "- Warranty for 12 months after go-live and comprehensive AMC for three years", "",
         "### Out of scope / provided by others", "",
         "- SDC infrastructure (servers, storage, network, backup media) — provided by the State Data Centre",
         "- Test centre premises, terminals, webcams, network and seating — arranged by KEONICS and centres",
         "- Licences and transaction charges of third-party services (payment gateway, SMS, Aadhaar, e-Sign)",
         "- Re-evaluation of answer scripts (not permitted under the RFP)", "",
         "### Service line items (quoted separately from software features)", "",
         "- Question content creation by Subject Matter Experts (at least 2,000 questions at launch, 5% growth per year)",
         "- Need-based offsite helpdesk with dedicated numbers",
         "- Annual Safe-to-Host audit by a CERT-In empanelled auditor and SSL certificates",
         "- Training, user manuals and how-to videos; exit management and knowledge transfer", "",
         "## 5. Users and Roles", "",
         "| Role | Responsibility |", "|------|----------------|",
         "| Candidate | Registers, books, pays, takes mock and main tests, downloads results and certificates |",
         "| Registration Centre Operator | Registers prospective candidates at designated centres and captures photographs |",
         "| Scrutiny Officer | Validates registrations and support requests |",
         "| SME Author / Reviewer / Technical Committee | Authors, reviews and approves questions |",
         "| Question Bank Administrator | Manages blueprint, banks and paper generation |",
         "| Proctor / Centre Operator | Runs sessions, validates candidates, handles incidents |",
         "| Examination Monitoring Unit | Monitors all centres live and takes control actions |",
         "| Result Officer / KEONICS Approver | Finalises results and approves certificate issue |",
         "| Finance Officer | Reconciles fees and approves refunds |",
         "| Helpdesk Executive | Logs calls and makes outbound calls |",
         "| KEONICS Administrator / Super Admin | Configures the system and manages users |", "",
         "## 6. Functional Scope by Module", ""]
    for m in MODULES:
        ms = feats(m)
        L += [f"### {m['nn']}. {m['name']}", "",
              f"**Type:** {m['type']} · **RFP reference:** {m['source'].split('(', 1)[-1].rstrip(')')}", "",
              f"**Purpose:** {m['purpose']}", "", "**Functional requirements:**", ""]
        L += [f"- {r}" for r in m["reqs"]]
        L += ["", f"**Acceptance criteria:** {m['accept']}", "", "**Key capabilities:**", "",
              "| Feature ID | Feature | Description |", "|------------|---------|-------------|"]
        L += [f"| {f['id']} | {f['name']} | {f['short']} |" for f in ms]
        L.append("")
    L += ["## 7. End-to-End Workflows", "",
          "1. **Employee:** OTP account → KGID registration with HRMS fetch → auto-accept or scrutiny → booking (free first attempt) → hall ticket → mock tests → centre validation → test → immediate score → result finalisation → signed certificate → HRMS update.",
          "2. **Prospective candidate:** registration at designated centre with eKYC and live photo → scrutiny → booking and fee payment → hall ticket → test → result → certificate.",
          "3. **Question bank:** blueprint → authoring → review → Technical Committee approval → bank → paper set per session → sealed → released at session start.",
          "4. **Session:** calendar → bookings → minimum batch check and forced reschedule → proctor assignment → readiness → validation → test with proctoring → wrap-up and archival.",
          "5. **Malpractice:** alert → proctor / EMU review → case → committee decision → result withheld or cancelled.", "",
          "## 8. Security, Audit and Compliance", "",
          "- AES-256 encryption of credentials, question bank, answer keys, responses and results; TLS in transit; asymmetric encryption with keys in the CeG HSM",
          "- Question component segregation so that only the author sees all components",
          "- Audit trails test centre-wise, IP-wise and candidate-wise in the Annexure 9 format; tamper-evident",
          "- Safe-to-Host certification before go-live and annually; ISO 27001 aligned controls; IPv6 readiness",
          "- DPDP Act compliance and Aadhaar Data Vault", "",
          "## 9. Non-Functional Requirements", "",
          "| Area | Requirement |", "|------|-------------|",
          "| Concurrency | At least 2,000 concurrent candidates per session with screen and webcam recording |",
          "| Response time | Under 1 second on LAN for test interactions |",
          "| Availability | 99.9% |",
          "| Compatibility | Device, operating system and browser agnostic |",
          "| Languages | Kannada and English |",
          "| Data location | SDC only; no data stored at centres after the session |",
          "| Backup | As per KEONICS policy; data live for the contract period |", "",
          "## 10. Legacy Data Porting", "",
          "Previous CLT candidates, attempts, results and certificates are ported so that attempt counts drive fee rules and earlier certificates remain verifiable. KEONICS provides data in agreed templates and signs off reconciliation before go-live.", "",
          "## 11. Delivery Phasing and Milestones", "",
          "| Milestone | Timeline | Payment |", "|-----------|----------|--------:|",
          "| Functional Requirement Specification | T + 1 month | 10% |",
          "| Design: UI/UX, MIS, technology stack, server requirements, data migration document | T + 1.5 months | 25% |",
          "| Development, data porting, UAT deployment, manuals and how-to videos | T + 2 months | — |",
          "| UAT completion | T + 2.5 months | 30% |",
          "| Safe-to-Host certification | T + 3 months | — |",
          "| Go-live (W) with 10 trouble-free sessions and source code handover | T + 3 months | 20% |",
          "| Warranty | W + 12 months | 5% |",
          "| Comprehensive AMC | 3 years after warranty | Quarterly |", "",
          "## 12. Deliverables", "",
          "- FRS, design documents, data migration document, test plans and reports",
          "- Application source code, binaries, mobile apps and deployment scripts",
          "- User manuals, operator and proctor manuals, how-to videos",
          "- Load test report, Safe-to-Host certificate, SLA reports", "",
          "## 13. Service Levels (summary of Annexure 4)", "",
          "- Availability 99.9% measured monthly; penalties per Annexure 4",
          "- Resolution times for critical, high and medium incidents per Annexure 4",
          "- Liquidated damages of 0.5% per week of delay up to 10%; performance bank guarantee of 5%", "",
          "## 14. Assumptions", "",
          "1. SDC provides compute, storage (including video recordings) and bandwidth sized per the capacity plan.",
          "2. HRMS look-up and result update APIs are provided by DPAR / CeG.",
          "3. Aadhaar AUA / KUA access, payment gateway, SMS (DLT) and e-mail accounts are provided by KEONICS / CeG.",
          "4. The CeG HSM partition and signing certificates are made available.",
          "5. Test centres provide terminals with webcams, required network bandwidth and permit installation of the secure exam client.",
          "6. Legacy CLT data is provided in exportable form.", "",
          "## 15. Appendix — RFP References", "",
          "- Terms of Reference (Section 4), Annexure 1 technical requirements, Special Conditions of Contract, Annexure 4 SLA, Annexure 9 audit log format", ""]
    (ROOT / f"SoW_{CODE}_Detailed.md").write_text("\n".join(L))


# ---------------------------------------------------------------- validation
REQUIRED = ("desc", "obj", "actors", "journey", "caps", "screens", "apis", "db", "tests", "docs", "inputs", "outputs",
            "rules", "valid", "calc", "flow", "search", "documents", "reports", "notif", "sec", "ints", "deps", "assume",
            "risks", "nfr", "est")


def validate():
    errors = []
    sow = (ROOT / f"SoW_{CODE}_Detailed.md").read_text()
    sow_feats = dict(re.findall(r"^\| (F\d{2}-\d{3}) \| ([^|]+?) \|", sow, re.M))
    ids = [f["id"] for f in FEATURES]
    if len(ids) != len(set(ids)):
        errors.append("duplicate feature IDs")
    for f in FEATURES:
        if f["id"] not in sow_feats:
            errors.append(f"{f['id']} missing in SOW")
        elif sow_feats[f["id"]].strip() != f["name"]:
            errors.append(f"{f['id']} name mismatch: SOW '{sow_feats[f['id']].strip()}' vs '{f['name']}'")
        if len(f.get("journey", [])) != 8:
            errors.append(f"{f['id']} journey has {len(f.get('journey', []))} points")
        for k in REQUIRED:
            if k not in f:
                errors.append(f"{f['id']} missing {k}")
        for area, key in AREA_KEYS:
            for name, cx in f.get(key, []):
                if cx not in CX:
                    errors.append(f"{f['id']} bad complexity {cx} on {name}")
    for fid in sow_feats:
        if fid not in ids:
            errors.append(f"{fid} in SOW but no feature data")
    for m in MODULES:
        if f"### {m['nn']}. {m['name']}" not in sow:
            errors.append(f"module heading missing in SOW: {m['nn']}. {m['name']}")
    if errors:
        raise SystemExit("Validation failed:\n" + "\n".join(errors))


# ---------------------------------------------------------------- Step 5
def render_feature(f):
    m = module_of(f)
    lines = [
        f"# {f['id']} — {f['name']}", "",
        "## Feature Specification", "",
        "| Field | Value |", "|-------|-------|",
        f"| Feature ID | {f['id']} |",
        f"| Feature Name | {f['name']} |",
        f"| Module | {m['nn']} — {m['name']} |",
        f"| Module Type | {m['type']} |",
        f"| Release | {f['notes']} |", "",
        "### Description", f["desc"], "",
        "### Business Objective", f["obj"], "",
        "### Actors", bullets(f["actors"]), "",
        "### User Journey",
    ]
    labels = ["Who creates", "How they create", "How they edit", "How they send for approval",
              "How the approver sees pending items", "How the approver decides",
              "How the applicant / maker sees history", "How the approver / auditor sees history"]
    for i, (lab, txt) in enumerate(zip(labels, f["journey"]), 1):
        lines.append(f"{i}. **{lab}:** {txt}")
    lines += ["", "### Functional Capabilities", bullets(f["caps"]), ""]
    lines += ["### UI Screens", bullets([s for s, _ in f["screens"]]) if f["screens"] else "- None (no user interface in this feature)", ""]
    lines += ["### Logical API Operations", bullets([a for a, _ in f["apis"]]) if f["apis"] else "- None", ""]
    for title, key in [("Inputs", "inputs"), ("Outputs", "outputs"), ("Business Rules", "rules"),
                       ("Validations", "valid"), ("Calculations", "calc")]:
        lines += [f"### {title}", bullets(f[key]), ""]
    lines += ["### Workflow", f["flow"], ""]
    for title, key in [("Search Requirements", "search"), ("Documents", "documents"), ("Reports", "reports"),
                       ("Notifications", "notif"), ("Security", "sec"), ("Integrations", "ints"),
                       ("Dependencies", "deps"), ("Assumptions", "assume"), ("Risks", "risks"),
                       ("Non Functional Requirements", "nfr")]:
        lines += [f"### {title}", bullets(f[key]), ""]
    lines += ["## Estimation Breakdown", "", "| Area | Artefact | Complexity | Hours |", "|------|----------|------------|------:|"]
    for area, name, cx, h in rows(f):
        lines.append(f"| {area} | {name} | {cx} | {h} |")
    lines += ["", f"**Feature Total: {f['hours']} hours**", ""]
    a = f["areas"]
    lines += ["| UI | API | Database | Integration | Testing | Documentation | Total |",
              "|---:|---:|---:|---:|---:|---:|---:|",
              f"| {a['UI']} | {a['API']} | {a['Database']} | {a['Integration']} | {a['Testing']} | {a['Documentation']} | **{f['hours']}** |",
              "", "### Estimation Assumptions",
              "- Every UI screen and logical API operation listed above is estimated as its own row using the Estimation Master.",
              "- Enterprise-grade implementation: validation, error handling, logging, audit hooks, security checks and unit tests included in each row."]
    lines += [f"- {e}" for e in f["est"]]
    if not f.get("integ"):
        lines.append("- No external integration channel in this feature; Integration row omitted.")
    return "\n".join(lines) + "\n"


def write_features():
    for m in MODULES:
        fdir = ROOT / m["folder"] / "features"
        fdir.mkdir(parents=True, exist_ok=True)
        for old in fdir.glob("F*.md"):
            old.unlink()
        ms = feats(m)
        for f in ms:
            f["file"] = f"{m['folder']}/features/{f['id']}-{f['slug']}.md"
            (ROOT / f["file"]).write_text(render_feature(f))
        m["hours"] = sum(f["hours"] for f in ms)
        idx = ["# Feature Index", "", f"## {m['nn']} — {m['name']}", "",
               "| Feature ID | Feature Name | Total Hours | Notes |", "|------------|--------------|------------:|-------|"]
        idx += [f"| {f['id']} | {f['name']} | {f['hours']} | {f['notes']} |" for f in ms]
        idx += ["", f"**Module Total:** **{m['hours']} hours**", ""]
        (fdir / "_index.md").write_text("\n".join(idx))
        readme = [f"# {m['nn']}. {m['name']}", "", f"Type: {m['type']}", f"Summary: {m['summary']}.", "",
                  f"- **Source:** {m['source']}", f"- **In Scope:** {m['in_scope']}", f"- **Out of Scope:** {m['out_scope']}",
                  f"- **Key Actors:** {m['actors']}", f"- **Dependencies:** {m['deps']}",
                  f"- **Notes:** {m['notes'] or '—'}", "",
                  f"Features: {len(ms)} · **{m['hours']} hours** (see `features/_index.md`)", ""]
        (ROOT / m["folder"] / "README.md").write_text("\n".join(readme))


# ---------------------------------------------------------------- Steps 2-3
def write_modules_md():
    out = [f"# {PROJECT} — Module List", "", f"Source: `SoW_{CODE}_Detailed.md` (derived from `261013-Call 3-RFP-CLT DC System-v1-0.pdf`, Section 6)", "",
           "| # | Module | Type | Summary | Source |", "|---|--------|------|---------|--------|"]
    for m in MODULES:
        src = "Derived" if m["source"].startswith("Derived") else "Explicit"
        out.append(f"| {m['nn']} | {m['name']} | {m['type']} | {m['summary']} | {src} |")
    out += ["", "---", "", "## Module Details", ""]
    for m in MODULES:
        out += [f"### {m['nn']}. {m['name']}", "", f"- **Type:** {m['type']}", f"- **Summary:** {m['summary']}.",
                f"- **Source:** {m['source']}", f"- **In Scope:** {m['in_scope']}", f"- **Out of Scope:** {m['out_scope']}",
                f"- **Key Actors:** {m['actors']}", f"- **Dependencies:** {m['deps']}", f"- **Notes:** {m['notes'] or '—'}",
                "- **Features:** " + ", ".join(f"{f['id']} {f['name']}" for f in feats(m)), ""]
    out += ["## Excluded / Embedded Capabilities", "",
            "- **Workflow / approvals:** embedded in registration scrutiny (F04-005), question review (F06-003), forced rescheduling (F05-006), malpractice (F10-005), results (F11-003) and certificate issue (F11-005); no separate workflow engine module.",
            "- **Document management:** embedded — photos and documents (F04-002 to F04-004), media (F06-002), hall tickets (F05-004), scans (F08-006), recordings (F10-006), certificates (F11-004).",
            "- **Scheduler / background jobs:** embedded in seat hold expiry (F05-001), minimum batch check (F05-006), paper generation (F07-001), reminders (F14-001), retention purge (F10-006), reconciliation (F05-003).",
            "- **Question content creation by SMEs:** service line item; the Content Authoring Tool and workflow are in Module 06.",
            "- **Helpdesk staffing, Safe-to-Host audit fee, SSL certificate, training and how-to videos, CAMC and exit management:** service or commercial line items outside feature hours.",
            "- **Re-evaluation:** not permitted under the RFP.", "",
            "## Assumptions", "",
            "1. Access model is User → Role → Permissions with district / centre data scope; candidates use a separate OTP login.",
            "2. HRMS, Aadhaar, payment gateway, HSM and DSC / e-Sign connectors are built in Module 15; business features call them.",
            "3. Full screen and webcam video recording is in go-live scope; storage and bandwidth are provided by SDC per the capacity plan.",
            "4. Field test and difficulty calibration (F06-005) is conditional on KEONICS requiring it.",
            "5. The typing test (Part 2, English and Kannada) is in go-live scope as per the RFP syllabus.",
            "6. KEONICS supplies templates, fee rules, centre data and legacy data in agreed formats.", "",
            "## Risks", "",
            "1. Go-live at T + 3 months with development complete by T + 2 months is tight for the size of the scope.",
            "2. Video recording at 2,000 concurrent candidates needs large storage and centre / SDC bandwidth.",
            "3. HRMS and Aadhaar AUA / KUA onboarding timelines are outside the Service Provider's control.",
            "4. Secure exam client lockdown depends on operating system versions at centres.",
            "5. Legacy CLT data quality is unknown.", ""]
    (ROOT / f"{CODE}_Modules.md").write_text("\n".join(out))


def totals():
    all_h = sum(f["hours"] for f in FEATURES)
    go = sum(f["hours"] for f in FEATURES if f["notes"] == GO_LIVE)
    return all_h, go, all_h - go


def write_registry():
    all_h, go, opt = totals()
    reg = dict(project_code=CODE, project_name=PROJECT, client=CLIENT, output_root="Projects/Keonics-Exam-Module/CLT_DC_System",
               source="261013-Call 3-RFP-CLT DC System-v1-0.pdf", sow=f"SoW_{CODE}_Detailed.md",
               updated_at=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
               totals=dict(all_hours=all_h, go_live_hours=go, conditional_hours=opt, feature_count=len(FEATURES),
                           module_count=len(MODULES)),
               modules=[])
    for m in MODULES:
        reg["modules"].append(dict(nn=m["nn"], name=m["name"], slug=m["slug"], type=m["type"], folder=m["folder"],
                                   feature_count=len(feats(m)), total_hours=m["hours"],
                                   features=[dict(id=f["id"], seq=f["seq"], name=f["name"], slug=f["slug"], file=f["file"],
                                                  hours=f["hours"], notes=f["notes"], areas=f["areas"]) for f in feats(m)]))
    (ROOT / f"{CODE}_Registry.json").write_text(json.dumps(reg, indent=2, ensure_ascii=False) + "\n")


# ---------------------------------------------------------------- Step 6
def area_sum(fs):
    t = {a: 0 for a, _ in AREA_KEYS}
    for f in fs:
        for a in t:
            t[a] += f["areas"][a]
    return t


def pct(x, base):
    return f"{(100 * x / base):.1f}%" if base else "0.0%"


def write_summaries():
    all_h, go, opt = totals()
    by_type = {}
    for m in MODULES:
        by_type[m["type"]] = by_type.get(m["type"], 0) + m["hours"]
    areas = area_sum(FEATURES)
    cond = [f for f in FEATURES if f["notes"] != GO_LIVE]
    top = sorted(FEATURES, key=lambda f: -f["hours"])[:10]
    build_weeks = 7  # design sign-off overlaps build; development complete by T + 2 months, UAT by T + 2.5 months

    s = [f"# {CODE} — Estimation Summary (Internal)", "", f"**Project:** {PROJECT}  ", f"**Client:** {CLIENT}  ",
         f"**Project code:** {CODE}  ", f"**As of:** {TODAY}  ",
         f"**Source:** `SoW_{CODE}_Detailed.md` · feature indexes under `modules/*/features/_index.md`", "",
         "## Totals", "", "| Metric | Value |", "|--------|------:|", f"| Modules | {len(MODULES)} |", f"| Features | {len(FEATURES)} |",
         f"| **Recommended go-live hours** | **{go}** |", f"| Conditional hours | {opt} |", f"| **All-in hours** | **{all_h}** |",
         f"| All-in person-days (8 h) | {all_h / 8:.0f} |", "",
         "Conditional: " + ", ".join(f"{f['id']} {f['name']} ({f['hours']} h)" for f in cond) + ".", "",
         "## Per-module hours", "", "| # | Module | Type | Features | Hours | % of all-in |", "|---|--------|------|---------:|------:|------------:|"]
    for m in MODULES:
        s.append(f"| {m['nn']} | {m['name']} | {m['type']} | {len(feats(m))} | {m['hours']} | {pct(m['hours'], all_h)} |")
    s += ["", "## Hours by module type", "", "| Type | Hours | % |", "|------|------:|---:|"]
    s += [f"| {t} | {h} | {pct(h, all_h)} |" for t, h in sorted(by_type.items())]
    s += ["", "## Hours by estimation area", "", "| Area | Hours | % |", "|------|------:|---:|"]
    s += [f"| {a} | {h} | {pct(h, all_h)} |" for a, h in areas.items()]
    s += ["", "## Top 10 features by effort", "", "| Feature | Name | Hours |", "|---------|------|------:|"]
    s += [f"| {f['id']} | {f['name']} | {f['hours']} |" for f in top]
    s += ["", "## Timeline check against the RFP", "",
          "The RFP expects FRS at T + 1 month, design at T + 1.5 months, development and UAT deployment at T + 2 months, UAT at "
          "T + 2.5 months and go-live with Safe-to-Host at T + 3 months. Allowing build to start once the FRS is drafted gives about "
          f"**{build_weeks} build weeks**. Go-live effort of **{go} hours** over {build_weeks} weeks at 40 hours per week needs roughly "
          f"**{go / (build_weeks * 40):.0f} full-time build-and-test engineers**, plus project management, design, DevOps and security. "
          "This is a high-risk timeline. Options: a larger team with parallel streams (candidate journey, question bank, exam engine and "
          "proctoring, integrations), reuse of an existing exam platform base, or agreement with KEONICS on a phased go-live "
          "(registration, booking and question bank first; full proctoring analytics, query tool and calibration after the first sessions).", "",
          "## Video recording sizing (go-live scope)", "",
          "Screen plus webcam recording at about 350 MB per candidate per 90-minute session gives about **700 GB per full session** of "
          "2,000 candidates. With 90-day retention and the expected session load, plan for several TB of SDC object storage and at "
          "least 50 Mbps per centre and 200 Mbps into SDC at peak. Storage and bandwidth are SDC provisions; the capacity plan (F17-002) "
          "fixes the figures before go-live.", "",
          "## Not included in feature hours", "",
          "- Project management, FRS and design documentation phase",
          "- Question content creation by SMEs (at least 2,000 questions at launch, 5% growth per year) and Technical Committee costs",
          "- Need-based offsite helpdesk staffing and dedicated numbers",
          "- Safe-to-Host audit fee (annual, CERT-In empanelled auditor) and SSL certificates",
          "- Training, user manuals beyond technical documentation, and how-to videos",
          "- Warranty (12 months) and comprehensive AMC (3 years)",
          "- Exit management and knowledge transfer services",
          "- SDC hosting, storage and bandwidth; payment gateway, SMS, Aadhaar and e-Sign transaction charges", "",
          "## Open assumptions / risks", "",
          "1. HRMS provides look-up by KGID and accepts result push; otherwise a file-based exchange is needed.",
          "2. Aadhaar authentication and eKYC through the state AUA / KUA via CeG; onboarding is outside our control.",
          "3. CeG HSM partition and client libraries are available for key management and signing.",
          "4. Centres permit installation of the secure exam client on Windows, Linux or macOS terminals with webcams.",
          "5. Kannada typing layout (Nudi / KGP or InScript) is fixed in the FRS.",
          "6. Field test and calibration (F06-005) is conditional and outside go-live hours.",
          "7. Legacy CLT data is provided in templates; cleansing is KEONICS' responsibility.",
          "8. The RFP contract period wording varies (3 years 3 months vs 2 + 2 years); confirm during pre-bid.", "",
          "## Suggested next actions", "",
          "1. Review modules, features and hours with the delivery lead.",
          "2. Decide team size and phasing against the T + 3 months go-live (see timeline check).",
          "3. Raise pre-bid queries on HRMS APIs, AUA / KUA, HSM, storage for video recordings and typing layout.",
          f"4. Share the client Word export: `{CODE}_Estimation_Summary_FeatureWise_Client.docx`.", ""]
    (ROOT / f"{CODE}_Estimation_Summary.md").write_text("\n".join(s))

    fw = [f"# {CODE} — Feature-wise Estimation Summary (Internal)", "", f"**Project:** {PROJECT}  ", f"**Project code:** {CODE}  ",
          f"**As of:** {TODAY}", "", "## Grand totals", "", "| Bucket | Hours |", "|--------|------:|",
          f"| Recommended go-live | {go} |", f"| Conditional | {opt} |", f"| All-in | {all_h} |",
          f"| Features | {len(FEATURES)} |", f"| Modules | {len(MODULES)} |", "",
          "## Module-wise summary", "", "| # | Module | Type | Features | UI | API | DB | Integ. | Test | Docs | Hours |",
          "|---|--------|------|---------:|---:|---:|---:|---:|---:|---:|------:|"]
    for m in MODULES:
        a = area_sum(feats(m))
        fw.append(f"| {m['nn']} | {m['name']} | {m['type']} | {len(feats(m))} | {a['UI']} | {a['API']} | {a['Database']} | "
                  f"{a['Integration']} | {a['Testing']} | {a['Documentation']} | {m['hours']} |")
    fw.append(f"| | **Total** | | **{len(FEATURES)}** | **{areas['UI']}** | **{areas['API']}** | **{areas['Database']}** | "
              f"**{areas['Integration']}** | **{areas['Testing']}** | **{areas['Documentation']}** | **{all_h}** |")
    fw += ["", "## Feature-wise detail", ""]
    for m in MODULES:
        fw += [f"### {m['nn']} — {m['name']}", "",
               "| Feature ID | Feature Name | Screens | APIs | UI | API | DB | Integ. | Test | Docs | Hours | Notes |",
               "|------------|--------------|--------:|-----:|---:|---:|---:|---:|---:|---:|------:|-------|"]
        for f in feats(m):
            a = f["areas"]
            fw.append(f"| {f['id']} | {f['name']} | {len(f['screens'])} | {len(f['apis'])} | {a['UI']} | {a['API']} | {a['Database']} | "
                      f"{a['Integration']} | {a['Testing']} | {a['Documentation']} | {f['hours']} | {f['notes']} |")
        fw += ["", f"**Module total: {m['hours']} hours**", ""]
    fw += ["## Effort profile", "", "| Area | Hours | % |", "|------|------:|---:|"]
    fw += [f"| {a} | {h} | {pct(h, all_h)} |" for a, h in areas.items()]
    fw += ["", "## Indicative phasing (from the SOW)", "",
           "| Phase | Timeline | Modules |", "|-------|----------|---------|",
           "| FRS | T to T + 1 month | All (requirements) |",
           "| Design and foundation | T + 0.5 to T + 1.5 months | 01, 02, 14, 15, 16, 17 |",
           "| Candidate journey and question bank | T + 1 to T + 2 months | 03, 04, 05, 06, 07 |",
           "| Exam engine, centre, proctoring, results | T + 1 to T + 2 months | 08, 09, 10, 11, 12 |",
           "| MIS and data porting | T + 1.5 to T + 2 months | 13, 18 |",
           "| UAT | T + 2 to T + 2.5 months | All |",
           "| Safe-to-Host and go-live | T + 2.5 to T + 3 months | All |", "",
           "## Exclusions", "",
           "- PM, FRS and design phase, SME question content, helpdesk staffing, Safe-to-Host audit fee and SSL, training and how-to videos, warranty and CAMC, exit management, SDC hosting and third-party transaction charges", ""]
    (ROOT / f"{CODE}_Estimation_Summary_FeatureWise.md").write_text("\n".join(fw))


# ---------------------------------------------------------------- DOCX helpers
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

NAVY = RGBColor(0x1B, 0x3A, 0x5F)


def base_doc():
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = NAVY
        doc.styles[name].font.name = "Calibri"
    return doc


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_runs(par, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = par.add_run(part.replace("`", ""))
        run.bold = i % 2 == 1


def add_table(doc, header, body, num_cols=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        shade(c, "1B3A5F")
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_runs(p, str(v))
            for r in p.runs:
                r.font.size = Pt(9.5)
            if i in num_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if ri % 2 == 1:
                shade(cells[i], "F2F4F7")
    doc.add_paragraph()
    return t


def md_to_docx(md_path: Path, out_path: Path):
    doc = base_doc()
    lines = md_path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            body = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                body.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, header, body)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2 if re.match(r"^\d{2}\. ", line[4:]) else 3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(20)
            r.font.color.rgb = NAVY
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip() in ("---", ""):
            pass
        else:
            add_runs(doc.add_paragraph(), line.rstrip(" "))
        i += 1
    doc.save(out_path)


# ---------------------------------------------------------------- Step 7
CLIENT_NOTES = {}


def write_client_docx():
    all_h, go, opt = totals()
    doc = base_doc()
    for text, size, bold in [("Module & Feature-wise Effort Estimate", 22, True), (PROJECT, 14, True),
                             (f"Prepared for {CLIENT} · Client Copy · Effort in person-hours", 11, False),
                             (f"Date: {TODAY}", 10, False)]:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("This document sets out the module and feature-wise effort for the Online Computer Literacy Test (CLT) and Digital "
                      "Certification System for KEONICS on behalf of the Center for E-Governance. The system covers the full CLT cycle: "
                      "public portal and mobile application, candidate registration with HRMS validation and Aadhaar eKYC, online test "
                      "booking with fees and hall tickets, the question bank and authoring tool, paper generation and mock tests, test "
                      "centre operations, the online examination with English and Kannada typing tests, screen and webcam video "
                      "proctoring, evaluation, digitally signed certificates, central monitoring, MIS and legacy data porting.")
    doc.add_paragraph("Hours cover the design and build of the application software, its integrations, security and audit controls, "
                      "performance engineering for 2,000 concurrent candidates, operational tooling and porting utilities, with associated "
                      "testing and technical documentation.")
    doc.add_paragraph("Hours do not cover project management, the FRS and design phase, question content creation, helpdesk staffing, the "
                      "Safe-to-Host audit, training and how-to videos, warranty and comprehensive AMC, or hosting and third-party service charges.")

    doc.add_heading("2. Summary Totals", level=1)
    add_table(doc, ["Bucket", "Person-hours"], [["Recommended for go-live", go], ["Conditional (taken up only if required by KEONICS)", opt],
                                                ["**Total (all items)**", f"**{all_h}**"], ["Modules", len(MODULES)],
                                                ["Features", len(FEATURES)]], num_cols=(1,))

    doc.add_heading("3. Module-wise Summary", level=1)
    add_table(doc, ["Module", "Type", "Features", "Hours", "% of go-live"],
              [[f"{m['nn']}. {m['name']}", m["type"], len(feats(m)), m["hours"],
                pct(sum(f["hours"] for f in feats(m) if f["notes"] == GO_LIVE), go)] for m in MODULES]
              + [["**Total**", "", f"**{len(FEATURES)}**", f"**{all_h}**", "**100%**"]], num_cols=(2, 3, 4))

    doc.add_heading("4. Feature-wise Detail", level=1)
    for m in MODULES:
        doc.add_heading(f"{m['nn']}. {m['name']}", level=2)
        doc.add_paragraph(m["summary"] + ".")
        add_table(doc, ["Feature ID", "Feature Name", "Hours", "Notes"],
                  [[f["id"], f["name"], f["hours"], f["client"]] for f in feats(m)]
                  + [["", "**Module total**", f"**{m['hours']}**", ""]], num_cols=(2,))

    doc.add_heading("5. Conditional Items and Commercial Line Items", level=1)
    cond = [f for f in FEATURES if f["notes"] != GO_LIVE]
    add_table(doc, ["Item", "Basis"],
              [[f"{f['id']} {f['name']}", f"{f['hours']} person-hours, taken up only if KEONICS requires field testing and calibration"] for f in cond]
              + [["Question content creation by Subject Matter Experts", "At least 2,000 bilingual questions at launch and 5% growth per year; quoted separately"],
                 ["Helpdesk with dedicated numbers", "Need-based offsite helpdesk; quoted separately"],
                 ["Safe-to-Host audit and SSL certificates", "Annual audit by a CERT-In empanelled auditor; at actuals"],
                 ["Comprehensive AMC for three years after warranty", "Quoted separately, payable quarterly"],
                 ["Change of scope", "At the rate quoted in the financial proposal"]])

    doc.add_heading("6. Effort Profile", level=1)
    areas = area_sum(FEATURES)
    label = {"UI": "User interface (screens and portal pages)", "API": "Application services and business rules",
             "Database": "Data model and storage", "Integration": "Integrations", "Testing": "Testing", "Documentation": "Documentation"}
    add_table(doc, ["Work area", "Hours", "% of total"], [[label[a], h, pct(h, all_h)] for a, h in areas.items()], num_cols=(1, 2))
    top = sorted(FEATURES, key=lambda f: -f["hours"])[:10]
    doc.add_paragraph("Largest features by effort:")
    add_table(doc, ["Feature ID", "Feature Name", "Hours"], [[f["id"], f["name"], f["hours"]] for f in top], num_cols=(2,))
    by_type = {}
    for m in MODULES:
        by_type[m["type"]] = by_type.get(m["type"], 0) + m["hours"]
    add_table(doc, ["Module type", "Hours", "% of total"], [[t, h, pct(h, all_h)] for t, h in sorted(by_type.items())], num_cols=(1, 2))

    doc.add_heading("7. Indicative Delivery Phasing", level=1)
    add_table(doc, ["Phase", "Timeline", "Scope"], [
        ["Functional Requirement Specification", "T + 1 month", "Requirements for all modules, integration positions, test pattern and rules"],
        ["Design and foundation", "T + 0.5 to T + 1.5 months", "UI/UX, MIS design, technology stack, server requirements, data migration document; access, configuration, notifications, integrations, security and environments"],
        ["Candidate journey and question bank", "T + 1 to T + 2 months", "Portal and mobile application, registration and eKYC, booking, fees and hall ticket, authoring tool, question bank and paper generation"],
        ["Examination and certification", "T + 1 to T + 2 months", "Test centre operations, secure exam client and examination engine, video proctoring, evaluation, certificates, central monitoring"],
        ["MIS and data porting", "T + 1.5 to T + 2 months", "Dashboards, reports, query tool, legacy CLT data porting"],
        ["User acceptance testing", "T + 2 to T + 2.5 months", "UAT at SDC, load testing at 2,000 concurrent candidates"],
        ["Safe-to-Host and go-live", "T + 2.5 to T + 3 months", "Safe-to-Host certification, go-live with ten trouble-free sessions, source code handover"],
    ])

    doc.add_heading("8. Items Not Included in Feature Hours", level=1)
    for t in ["Project management and the FRS and design documentation phase",
              "Question content creation by Subject Matter Experts and Technical Committee costs",
              "Helpdesk staffing and dedicated helpdesk numbers",
              "Safe-to-Host audit by a CERT-In empanelled auditor and SSL certificates",
              "Training, user manuals and how-to videos",
              "Warranty (12 months) and comprehensive AMC (three years)",
              "Exit management and knowledge transfer",
              "SDC hosting, storage and bandwidth; payment gateway, SMS, e-mail, Aadhaar and e-Sign charges",
              "Test centre premises, terminals, webcams and network"]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("9. Key Planning Assumptions", level=1)
    for t in ["HRMS provides employee look-up by KGID and accepts the update of results and certificates.",
              "Aadhaar authentication and eKYC are available through the state AUA / KUA arrangement.",
              "The Center for E-Governance HSM and signing certificates are available for key management and digital signatures.",
              "KEONICS provides the payment gateway, SMS (with DLT registration) and e-mail accounts.",
              "SDC provides compute, storage for screen and webcam recordings and bandwidth as per the agreed capacity plan.",
              "Test centres provide terminals with webcams and network, and permit installation of the secure exam client.",
              "The English and Kannada typing tests form Part 2 of the examination; the Kannada keyboard layout is confirmed in the FRS.",
              "Legacy CLT data is provided by KEONICS in the agreed templates.",
              "Access is role-based with district and centre data scope; candidates use a separate OTP login."]:
        doc.add_paragraph(t, style="List Number")

    doc.add_heading("10. Disclaimer", level=1)
    doc.add_paragraph("Estimates are indicative for planning and commercial discussion. Final effort may vary with the HRMS, Aadhaar, "
                      "HSM and signing positions confirmed in the FRS, the proctoring and recording parameters, centre terminal "
                      "configurations, the volume and quality of legacy data, and any scope changes agreed with KEONICS.")
    doc.save(ROOT / f"{CODE}_Estimation_Summary_FeatureWise_Client.docx")


def main():
    write_sow()
    validate()
    write_features()
    write_modules_md()
    write_registry()
    write_summaries()
    md_to_docx(ROOT / f"SoW_{CODE}_Detailed.md", ROOT / f"SoW_{CODE}_Detailed.docx")
    write_client_docx()
    all_h, go, opt = totals()
    print(f"Features {len(FEATURES)} · go-live {go} · conditional {opt} · all-in {all_h}")
    for m in MODULES:
        print(f"  {m['nn']} {m['name']}: {len(feats(m))} features, {m['hours']} h")


if __name__ == "__main__":
    main()
